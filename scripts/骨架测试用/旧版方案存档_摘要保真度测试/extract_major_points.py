# -*- coding: utf-8 -*-
"""从节期纲目 docx 抽取「壹贰叁……」大点，纯本地解析，不调 API。

用法：
  python extract_major_points.py --trial   # 只跑 1真理启示 编号最小的一篇，打印大点
  python extract_major_points.py           # 跑全部 30 篇，写出 大点标注.xlsx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

SCRIPT_DIR = Path(__file__).resolve().parent
INPUT_ROOT = SCRIPT_DIR / "第一次测试" / "30个特会题目"
OUTPUT_XLSX = SCRIPT_DIR / "第一次测试" / "大点标注.xlsx"

FOLDER_TO_CATEGORY = {
    "1真理启示": "真理类",
    "2生命经历": "生命类",
    "3应用实行": "实行类",
}
FOLDER_ORDER = ("1真理启示", "2生命经历", "3应用实行")

# 大点：壹～拾。分隔可以是顿号/点/空白，也可以直接接下文（如「壹撒拉…」）。
# 不匹配「一、二、三」中点；也不把「十一」类中点当成大点。
MAJOR_MARKERS = ("壹", "贰", "叁", "肆", "伍", "陆", "柒", "捌", "玖", "拾")
MAJOR_RE = re.compile(
    r"^(?P<seq>壹|贰|叁|肆|伍|陆|柒|捌|玖|拾)(?:[、．.]\s*|(?=\s)|(?=[^\s、．.一二三四五六七八九十0-9])|$)"
)


def _numeric_prefix(path: Path) -> int:
    m = re.match(r"^(\d+)", path.stem)
    return int(m.group(1)) if m else 10**9


def list_docx_in_folder(folder: Path) -> list[Path]:
    files = [p for p in folder.glob("*.docx") if not p.name.startswith("~$")]
    return sorted(files, key=_numeric_prefix)


def read_nonempty_paras(path: Path) -> list[str]:
    doc = Document(str(path))
    return [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]


def extract_major_points(paras: list[str]) -> list[tuple[str, str]]:
    """返回 [(序号, 该段完整文字), ...]。"""
    hits: list[tuple[str, str]] = []
    for text in paras:
        m = MAJOR_RE.match(text)
        if not m:
            continue
        hits.append((m.group("seq"), text))
    return hits


def collect_jobs(*, trial: bool) -> list[tuple[Path, str]]:
    jobs: list[tuple[Path, str]] = []
    for folder_name in FOLDER_ORDER:
        folder = INPUT_ROOT / folder_name
        if not folder.is_dir():
            raise SystemExit(f"找不到输入目录: {folder}")
        category = FOLDER_TO_CATEGORY[folder_name]
        files = list_docx_in_folder(folder)
        if not files:
            raise SystemExit(f"目录无 docx: {folder}")
        if trial:
            if folder_name == "1真理启示":
                return [(files[0], category)]
            continue
        jobs.extend((p, category) for p in files)
    return jobs


def write_xlsx(rows: list[dict], skipped: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "大点"
    headers = ["编号", "分类", "主题", "大点序号", "大点内容", "人工标注"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append(
            [
                row["编号"],
                row["分类"],
                row["主题"],
                row["大点序号"],
                row["大点内容"],
                "",
            ]
        )
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=6):
        for c in r:
            c.alignment = wrap
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 80
    ws.column_dimensions["F"].width = 18

    ws2 = wb.create_sheet("解析失败")
    ws2.append(["编号", "分类", "主题", "文件路径", "原因"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for item in skipped:
        ws2.append(
            [
                item["编号"],
                item["分类"],
                item["主题"],
                item["路径"],
                item["原因"],
            ]
        )
    ws2.column_dimensions["A"].width = 36
    ws2.column_dimensions["C"].width = 40
    ws2.column_dimensions["D"].width = 60
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="抽取节期纲目大点")
    parser.add_argument(
        "--trial",
        action="store_true",
        help="只处理 1真理启示 中编号最小的一篇并打印，不写 Excel",
    )
    args = parser.parse_args()

    jobs = collect_jobs(trial=args.trial)
    rows: list[dict] = []
    skipped: list[dict] = []

    for i, (path, category) in enumerate(jobs, 1):
        print(f"[{i}/{len(jobs)}] {category} {path.name}", flush=True)
        paras = read_nonempty_paras(path)
        title = paras[0] if paras else ""
        points = extract_major_points(paras)
        if not points:
            skipped.append(
                {
                    "编号": path.stem,
                    "分类": category,
                    "主题": title,
                    "路径": str(path),
                    "原因": "未找到壹贰叁……格式的大点",
                }
            )
            print("  -> 未找到大点，已记录并跳过")
            continue
        if args.trial:
            print(f"主题: {title}")
            print(f"大点数: {len(points)}")
            print("---")
            for seq, text in points:
                print(f"[{seq}] {text}")
                print("---")
        for seq, text in points:
            rows.append(
                {
                    "编号": path.stem,
                    "分类": category,
                    "主题": title,
                    "大点序号": seq,
                    "大点内容": text,
                }
            )

    if skipped and not args.trial:
        print("解析失败的文件：")
        for item in skipped:
            print(f"  - {item['分类']} {item['编号']}：{item['原因']}")

    if not args.trial:
        write_xlsx(rows, skipped, OUTPUT_XLSX)
        print(
            f"已写入 {OUTPUT_XLSX}（大点 {len(rows)} 条，失败文件 {len(skipped)} 个）"
        )


if __name__ == "__main__":
    main()
