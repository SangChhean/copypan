# -*- coding: utf-8 -*-
"""为 29 篇节期纲目生成评审 docx：标题 / 负担说明 / 生成的龙骨 / 龙骨评价。

用法：
  python generate_review_docx.py --id "2021秋长-1"   # 只生成一篇
  python generate_review_docx.py --folder 2021秋长
  python generate_review_docx.py                      # 生成全部
"""
from __future__ import annotations

import argparse
import json
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from dotenv import load_dotenv

from generate_skeleton import SIX_STAGE_CRITERIA

SCRIPT_DIR = Path(__file__).resolve().parent
BACKEND_DIR = SCRIPT_DIR.parents[1] / "back_mic" / "backend"
BURDEN_JSON = SCRIPT_DIR / "第三次测试" / "burden_all29_八稿.json"
SKELETON_JSON = SCRIPT_DIR / "第三次测试" / "skeleton_all29_八稿.json"
SRC_ROOT = SCRIPT_DIR
OUT_ROOT = SCRIPT_DIR / "第三次测试" / "评审docx"

CATEGORY_TO_FOLDER = {
    "2021秋长": "2021秋长",
    "2024国殇": "2024国殇",
    "2026国殇": "2026国殇",
    "2026秋长": "2026秋长",
}

MODEL = "claude-opus-5"
MAX_TOKENS = 8192

SYSTEM_PROMPT = """你是一位资深的职事信息研究者，具备纲目写作与神学内容分析的实务经验。你的任务是站在独立评审的角度，客观、平衡地评价另一位研究者对一篇负担说明所做的「六阶段」骨架划分是否合理——不预设这份划分一定正确，也不因为它出自资深研究者之手就照单全收，该指出问题就指出问题。
严格只输出最终评价报告，不输出内部思考步骤、分析过程、分隔线。"""

USER_PROMPT_TEMPLATE = """你的任务是：结合「原纲目全文」（这是负担说明和龙骨划分真正的原始依据，包含负担说明和龙骨可能没有覆盖到的中点、小点细节），评价下面这份「龙骨结果」的六阶段划分是否合理。

━━━━━━━━━━━━━━━━━━━
【六阶段判断依据】（评价时依据这套标准，不要自创别的标准）
━━━━━━━━━━━━━━━━━━━

""" + SIX_STAGE_CRITERIA + """

━━━━━━━━━━━━━━━━━━━
【用户输入】
━━━━━━━━━━━━━━━━━━━
- 主题：{主题}
- 原纲目全文：{原纲目全文}
- 负担说明：{负担说明}
- 待评价的龙骨结果：{龙骨结果}

━━━━━━━━━━━━━━━━━━━
【评价要求】
━━━━━━━━━━━━━━━━━━━

全篇评价控制在 400–600 字（汉字计数，含三个小标题）。第一段按六项逐项写，每项 1–2 句，不要铺开；骨架中未出现的项只写「未出现」，不展开。不要核对分类定义是否准确，不要写「遗漏核对」，不要把问题归因于负担说明是否压缩过度。

一、逐项评价六阶段
逐项评价六阶段（圣经根据、真理启示、主观经历、团体建造、应用实行、确定目标），每项 1–2 句。评价时不是去核对这一项符不符合分类定义——分类是否准确已经在别处处理过，这里要做的是凭自己的判断，像一个真正在读这篇纲目的人一样，说出这一段文字本身写得好不好、有没有力量，可以有自己的观点，不必字字都对照标准。

若这一项在骨架中出现：
- 评价它作为纲目这一段落，本身讲得扎不扎实：
  - 圣经根据：这处经文，读起来站不站得住、有没有说服力，是不是这篇论述扎实的根基，而不是硬凑上去的
  - 真理启示：这条真理讲得够不够深、说得够不够透，还是含糊笼统、点到为止
  - 主观经历：这段经历读起来真不真切、有没有打动人，还是流于表面的说法
  - 团体建造：这段团体的实际，写得具不具体、有没有画面感，还是空泛地提一句「身体」「召会」
  - 应用实行：这条路径听起来能不能真的照着做，还是原则性的空话
  - 确定目标：这个终极指向，读起来有没有力度、能不能让人心里一震，还是平平淡淡带过
- 如果这一项底下有好几条内容，它们彼此之间是不是有层层递进的关系，还是平行罗列
- 这一项的内容量，撑不撑得起一个纲目大点

六项都评完后，用 1–2 句给一个总体判断：整体推进顺不顺、读起来像不像一份可以直接拿去用的骨架、有没有一气呵成的感觉。

二、与原纲目的差异
只提有分量的差异，每处一行客观说明差异是什么，不需要判断这个差异是负担说明造成的还是龙骨判断造成的。最多写两三处。若无明显差异，写「与原纲目无有分量的差异」。

三、综合评价
一两句定性判断（例如「整体站得住，个别几处存疑」或「结构性偏差较大，建议重新检视」）。只针对这份龙骨本身下判断，不涉及负担说明是否压缩过度这类上游归因。不需要打分数。

━━━━━━━━━━━━━━━━━━━
【输出格式】
━━━━━━━━━━━━━━━━━━━
按以下三个部分输出，用平实的书面语写成一份评审报告，不需要输出 JSON 或其他结构化格式：

一、逐项评价六阶段
[...]

二、与原纲目的差异
[...]

三、综合评价
[...]"""


def _load_api_key() -> str:
    load_dotenv(BACKEND_DIR / ".env")
    key = (os.getenv("CLAUDE_API_KEY") or "").strip()
    if not key:
        raise SystemExit("未找到 CLAUDE_API_KEY（请检查 back_mic/backend/.env）")
    return key


OPTS: dict = {
    "title_from_id": False,
    "section2_title": "二、生成的龙骨",
    "flat_out": False,
    "review_json": None,
}

EVAL_HEAD_RE = re.compile(r"^[一二三]、\S")
CIRCLED_NUMS = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
SKELETON_HEAD_RE = re.compile(
    rf"^([{CIRCLED_NUMS}])\s*(?:[\[【](?P<br>[^】\]]+)[\]】]|(?P<label>[^\n]+))?\s*(?P<rest>.*)$",
    re.DOTALL,
)
STAGE_LABELS = ["圣经根据", "真理启示", "主观经历", "团体建造", "应用实行", "确定目标"]
STAGE_COLORS = {
    "圣经根据": "D9E2F3",
    "真理启示": "E4DFEC",
    "主观经历": "FCE4EC",
    "团体建造": "E2EFDA",
    "应用实行": "FFF2CC",
    "确定目标": "FCE4D6",
}
HEADER_FILL = "D9D9D9"


def ensure_out_dirs() -> None:
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    if OPTS["flat_out"]:
        return
    for folder in CATEGORY_TO_FOLDER.values():
        (OUT_ROOT / folder).mkdir(parents=True, exist_ok=True)


def load_json_list(path: Path) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise SystemExit(f"{path.name} 不是列表")
    return data


def stem_to_record_id(feast: str, stem: str) -> str:
    """msg. 1 标题 / 1标题 → {feast}-1标题"""
    s = re.sub(r"^msg\.\s*", "", stem, flags=re.I)
    s = re.sub(r"^(\d+)\s+", r"\1", s)
    return f"{feast}-{s}"


def find_source_docx(rec_id: str, category: str) -> Path:
    folder_name = CATEGORY_TO_FOLDER.get(category) or category
    folder = SRC_ROOT / folder_name
    exact = folder / f"{rec_id}.docx"
    if exact.is_file():
        return exact
    flat = SRC_ROOT / f"{rec_id}.docx"
    if flat.is_file():
        return flat
    search_dirs = []
    if folder.is_dir():
        search_dirs.append((folder, folder_name))
    else:
        for name in CATEGORY_TO_FOLDER.values():
            d = SRC_ROOT / name
            if d.is_dir():
                search_dirs.append((d, name))
    for directory, feast in search_dirs:
        for docx in directory.glob("*.docx"):
            if docx.name.startswith("~$"):
                continue
            if stem_to_record_id(feast, docx.stem) == rec_id:
                return docx
    raise FileNotFoundError(f"找不到原始纲目: {rec_id}（目录 {folder}）")


def read_full_outline(path: Path) -> str:
    doc = Document(str(path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    if not paras:
        raise ValueError(f"空文档: {path}")
    return "\n".join(paras)


def build_eval_prompt(topic: str, full_text: str, burden: str, skeleton: str) -> str:
    return (
        USER_PROMPT_TEMPLATE.replace("{主题}", topic)
        .replace("{原纲目全文}", full_text)
        .replace("{负担说明}", burden)
        .replace("{龙骨结果}", skeleton)
    )


def call_claude(api_key: str, user_prompt: str) -> str:
    import anthropic

    client = anthropic.Anthropic(api_key=api_key, timeout=600.0)
    kwargs = {
        "model": MODEL,
        "max_tokens": MAX_TOKENS,
        "system": SYSTEM_PROMPT,
        "messages": [{"role": "user", "content": user_prompt}],
    }
    last_empty_info = ""
    for attempt in range(2):
        if attempt:
            kwargs["max_tokens"] = max(int(kwargs["max_tokens"]), 16000)
        msg = client.messages.create(**kwargs)
        print(
            f"  API model={getattr(msg, 'model', MODEL)} "
            f"stop={getattr(msg, 'stop_reason', None)} attempt={attempt + 1}",
            flush=True,
        )
        parts: list[str] = []
        block_types: list[str] = []
        for block in msg.content or []:
            btype = getattr(block, "type", None) or type(block).__name__
            block_types.append(str(btype))
            if btype and btype != "text":
                continue
            text = getattr(block, "text", None)
            if text and str(text).strip():
                parts.append(str(text).strip())
        raw = "\n".join(parts).strip()
        if raw:
            return raw
        last_empty_info = (
            f"stop_reason={getattr(msg, 'stop_reason', None)} "
            f"block_types={block_types}"
        )
        print(f"  空正文，将重试 | {last_empty_info}", flush=True)
    raise RuntimeError(f"Claude 返回为空 | {last_empty_info}")


def _set_run_font(run, *, name: str, size_pt: float, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def add_doc_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(16)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    _set_run_font(run, name="黑体", size_pt=16, bold=True)


def add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    _set_run_font(run, name="黑体", size_pt=14, bold=True)


def add_body_paragraphs(doc: Document, text: str, *, italic: bool = False) -> None:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text or "")
    chunks = [c.strip() for c in re.split(r"\n+", text or "") if c.strip()]
    if not chunks:
        chunks = ["（空）"]
    for chunk in chunks:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.first_line_indent = Cm(0)
        is_eval_head = bool(EVAL_HEAD_RE.match(chunk))
        run = p.add_run(chunk)
        if is_eval_head:
            _set_run_font(run, name="黑体", size_pt=12, bold=True)
            p.paragraph_format.space_before = Pt(8)
        else:
            _set_run_font(run, name="宋体", size_pt=12)
            run.italic = italic


def _match_stage(label: str) -> str:
    for lab in STAGE_LABELS:
        if lab in (label or ""):
            return lab
    return ""


def parse_skeleton_rows(skeleton: str) -> list[tuple[str, str, str]]:
    """按空行分块，只取块首圈码+阶段名+内容，丢掉依据。返回 [(左列, 内容, 阶段名), ...]。"""
    rows: list[tuple[str, str, str]] = []
    for block in re.split(r"\n\s*\n", skeleton or ""):
        block = block.strip()
        if not block:
            continue
        m = SKELETON_HEAD_RE.match(block)
        if not m:
            continue
        num = m.group(1)
        label = (m.group("br") or m.group("label") or "").strip()
        rest = (m.group("rest") or "").strip()
        if label.startswith("内容"):
            rest = f"{label}\n{rest}".strip()
            label = ""
        body = re.sub(r"^内容[：:]\s*", "", rest).strip()
        body = re.split(r"\n依据[：:]", body, maxsplit=1)[0].strip()
        body = re.sub(r"[ \t]*\n+[ \t]*", "", body)
        stage = _match_stage(label)
        left = f"{num}{stage}" if stage else f"{num}{label}".strip()
        rows.append((left, body or "（空）", stage))
    return rows


def _shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _set_cell_width(cell, width_cm: float) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    twips = str(int(width_cm * 567))
    tcW.set(qn("w:w"), twips)
    tcW.set(qn("w:type"), "dxa")


def _set_cell_margins(cell, dxa: int = 80) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcMar"))
    if old is not None:
        tcPr.remove(old)
    tcMar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(dxa))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_table_col_widths(table, widths_cm: list[float]) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = str(int(sum(widths_cm) * 567))
    tblW.set(qn("w:w"), total)
    tblW.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 567)))
        grid.append(gc)
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            _set_cell_width(cell, w)
            _set_cell_margins(cell)


def _fill_cell_text(cell, text: str, *, header: bool = False, center: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if center or header:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    if header:
        _set_run_font(run, name="黑体", size_pt=11, bold=True)
    else:
        _set_run_font(run, name="宋体", size_pt=10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if header else WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_skeleton_table(doc: Document, skeleton: str) -> None:
    rows = parse_skeleton_rows(skeleton)
    table = doc.add_table(rows=1 + max(len(rows), 1), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left_w, right_w = 3.0, 14.4
    _set_table_col_widths(table, [left_w, right_w])

    _fill_cell_text(table.cell(0, 0), "编号", header=True)
    _fill_cell_text(table.cell(0, 1), "内容", header=True)
    _shade_cell(table.cell(0, 0), HEADER_FILL)
    _shade_cell(table.cell(0, 1), HEADER_FILL)

    if not rows:
        _fill_cell_text(table.cell(1, 0), "—")
        _fill_cell_text(table.cell(1, 1), "（空）")
        return
    for i, (left, body, stage) in enumerate(rows, 1):
        fill = STAGE_COLORS.get(stage, "F2F2F2")
        _fill_cell_text(table.cell(i, 0), left, center=True)
        _fill_cell_text(table.cell(i, 1), body)
        _shade_cell(table.cell(i, 0), fill)
        _shade_cell(table.cell(i, 1), fill)


def write_review_docx(
    out_path: Path,
    title: str,
    burden: str,
    evaluation: str,
    skeleton: str,
) -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    add_doc_title(doc, title)

    add_section_heading(doc, "一、负担说明")
    add_body_paragraphs(doc, burden)

    add_section_heading(doc, OPTS["section2_title"])
    add_skeleton_table(doc, skeleton)

    add_section_heading(doc, "三、龙骨评价")
    add_body_paragraphs(doc, evaluation)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(out_path))
        return out_path
    except PermissionError:
        alt = out_path.with_name(out_path.stem + "_新版.docx")
        doc.save(str(alt))
        print(f"  原文件被占用，已写入 {alt.name}", flush=True)
        return alt


def merge_records() -> list[dict]:
    skeletons = {x["id"]: x for x in load_json_list(SKELETON_JSON)}
    missing = []
    ordered = load_json_list(BURDEN_JSON)
    out = []
    for item in ordered:
        rec_id = item["id"]
        if rec_id not in skeletons:
            missing.append(f"缺龙骨: {rec_id}")
            continue
        out.append(
            {
                "id": rec_id,
                "分类": item.get("分类") or "",
                "主题": item.get("主题") or rec_id,
                "负担说明": item.get("负担说明") or "",
                "龙骨结果": skeletons[rec_id].get("龙骨结果") or "",
            }
        )
    if missing:
        raise SystemExit("数据对不齐：\n" + "\n".join(missing))
    return out


def process_one(rec: dict, api_key: str) -> tuple[Path, str]:
    rec_id = rec["id"]
    category = rec["分类"]
    folder = CATEGORY_TO_FOLDER.get(category, category)
    src = find_source_docx(rec_id, category)
    full_text = read_full_outline(src)
    prompt = build_eval_prompt(rec["主题"] or rec_id, full_text, rec["负担说明"], rec["龙骨结果"])
    evaluation = call_claude(api_key, prompt)
    if OPTS["flat_out"] or not folder:
        out_path = OUT_ROOT / f"{rec_id}.docx"
    else:
        out_path = OUT_ROOT / folder / f"{rec_id}.docx"
    title = rec_id if OPTS["title_from_id"] else (rec["主题"] or rec_id)
    written = write_review_docx(
        out_path,
        title,
        rec["负担说明"],
        evaluation,
        rec["龙骨结果"],
    )
    return written, evaluation


def _write_review_json(path: Path, items: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(items, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    global BURDEN_JSON, SKELETON_JSON, SRC_ROOT, OUT_ROOT
    parser = argparse.ArgumentParser(description="生成评审 docx")
    parser.add_argument("--trial", action="store_true", help="只生成第一篇")
    parser.add_argument("--folder", default="", help="只生成某个节期文件夹，如 2021秋长")
    parser.add_argument("--limit", type=int, default=0, help="在已筛选范围内只生成前 N 篇")
    parser.add_argument("--skip", type=int, default=0, help="在已筛选范围内跳过前 N 篇")
    parser.add_argument("--id", default="", help="只生成 id 包含该字符串的条目")
    parser.add_argument("--exclude", default="", help="跳过 id 包含该字符串的条目")
    parser.add_argument("--burden", default="", help="负担说明 JSON 路径")
    parser.add_argument("--skeleton", default="", help="龙骨结果 JSON 路径")
    parser.add_argument("--src-root", default="", help="原始纲目 docx 根目录（其下为 2021秋长 等）")
    parser.add_argument("--out-root", default="", help="评审 docx 输出目录")
    parser.add_argument("--review-json", default="", help="评价文本 JSON 输出路径")
    parser.add_argument("--title-from-id", action="store_true", help="文档标题用原始文件名（id）")
    parser.add_argument("--section2", default="", help="第二节标题，默认「二、生成的龙骨」")
    parser.add_argument("--flat-out", action="store_true", help="docx 直接写到 out-root，不分子文件夹")
    args = parser.parse_args()

    if args.burden:
        BURDEN_JSON = Path(args.burden)
    if args.skeleton:
        SKELETON_JSON = Path(args.skeleton)
    if args.src_root:
        SRC_ROOT = Path(args.src_root)
    if args.out_root:
        OUT_ROOT = Path(args.out_root)
    OPTS["title_from_id"] = bool(args.title_from_id)
    OPTS["section2_title"] = args.section2 or "二、生成的龙骨"
    OPTS["flat_out"] = bool(args.flat_out)
    OPTS["review_json"] = Path(args.review_json) if args.review_json else None

    ensure_out_dirs()
    records = merge_records()
    folder_to_cat = {v: k for k, v in CATEGORY_TO_FOLDER.items()}
    if args.id:
        jobs = [r for r in records if args.id in r["id"]]
        if not jobs:
            raise SystemExit(f"找不到 id 含「{args.id}」的条目")
    elif args.trial:
        jobs = records[:1]
        if not jobs:
            raise SystemExit("没有可生成的记录")
    elif args.folder:
        cat = folder_to_cat.get(args.folder) or args.folder
        jobs = [r for r in records if r["分类"] == cat]
        if not jobs:
            raise SystemExit(f"{args.folder} 下没有可生成的记录")
    else:
        jobs = records
    if args.exclude:
        jobs = [r for r in jobs if args.exclude not in r["id"]]
        if not jobs:
            raise SystemExit("排除后没有可生成的记录")
    if args.skip and args.skip > 0:
        jobs = jobs[args.skip :]
    if args.limit and args.limit > 0:
        jobs = jobs[: args.limit]

    api_key = _load_api_key()
    errors: list[str] = []
    reviews: list[dict] = []
    review_json = OPTS["review_json"]
    for i, rec in enumerate(jobs, 1):
        print(f"[{i}/{len(jobs)}] {rec['分类']} {rec['id']}", flush=True)
        try:
            out_path, evaluation = process_one(rec, api_key)
            reviews.append(
                {
                    "id": rec["id"],
                    "分类": rec.get("分类") or "",
                    "主题": rec.get("主题") or rec["id"],
                    "评价": evaluation,
                }
            )
            if review_json:
                _write_review_json(review_json, reviews)
            print(f"  已写入 {out_path}", flush=True)
        except Exception as e:
            errors.append(f"{rec['id']}: {e}")
            print(f"  失败: {e}", flush=True)
            continue

    if review_json:
        _write_review_json(review_json, reviews)
        print(f"评价 JSON 已写入 {review_json}（{len(reviews)} 篇）", flush=True)

    if errors:
        print("失败列表：")
        for line in errors:
            print(f"  - {line}")


if __name__ == "__main__":
    main()
