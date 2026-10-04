# -*- coding: utf-8 -*-
"""SC 三种独立 docx 导出验证。不调用真实 API。写入 logs/sc_export3.log。

用法（在 back_mic/backend）:
    python -m panai4.sc.verify_export3
"""
from __future__ import annotations

import io
import sys
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from panai4.export_docx import content_disposition
from panai4.sc.export_sc import (
    build_kind_docx,
    build_sc_docx,
    extract_burden_export,
    kind_docx_filename,
)

_REPO = Path(__file__).resolve().parents[4]
_LOG = _REPO / "logs" / "sc_export3.log"
_lines: list[str] = []

SAMPLE_SELECT = """【最内在】
生命｜依据：雅各变化的内里实际
变化｜依据：主题核心
雅各｜依据：人物本身
神人｜依据：变化的目标

【内在】
摔跤｜依据：毗努伊勒经历
祝福｜依据：抓住神得福
以色列｜依据：新名
天然生命｜依据：被对付的对象

【外在】
伯特利｜依据：回转之地
毗努伊勒｜依据：摔跤地点
拉结｜依据：家眷背景
以扫｜依据：弟兄关系
"""

SAMPLE_FULL = """【职事界定】
这是职事界定的内容，不应出现在导出正文。

【真理脉络】
这是真理脉络的内容，也不应出现。

【负担说明】
[360字左右]
第一行负担正文。
第二行负担正文。
第三行负担正文。
"""

SAMPLE_NO_BURDEN = """只有一段话。
没有负担说明标题。
"""


def log(msg: str = "") -> None:
    _lines.append(msg)
    print(msg)


def check(cond: bool, ok: str, fail: str) -> int:
    if cond:
        log(f"PASS: {ok}")
        return 0
    log(f"FAIL: {fail}")
    return 1


def _paras(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs]


def main() -> int:
    fails = 0
    log("=== SC EXPORT3 VERIFY ===")

    topic = "雅各的变化"
    log("--- 1 选词结果导出 ---")
    data, warn = build_kind_docx("select", topic, SAMPLE_SELECT)
    fails += check(warn is None, "选词无 warning", f"unexpected warning={warn}")
    doc = Document(io.BytesIO(data))
    paras = _paras(doc)
    title = doc.paragraphs[0]
    fails += check(title.text == topic, f"首段主题={title.text!r}", f"首段={title.text!r}")
    fails += check(
        title.alignment == WD_ALIGN_PARAGRAPH.CENTER,
        "首段居中",
        f"align={title.alignment}",
    )
    fails += check(bool(title.runs) and title.runs[0].bold, "首段加粗", "首段未加粗")
    size = title.runs[0].font.size
    fails += check(size == Pt(16), f"字号 16pt ({size})", f"字号={size}")
    expected_lines = SAMPLE_SELECT.replace("\r\n", "\n").split("\n")
    body = paras[1:]
    fails += check(
        body == expected_lines,
        f"正文 {len(body)} 段与示例逐行一致",
        f"mismatch body={body!r}\n expected={expected_lines!r}",
    )
    joined = "\n".join(paras)
    fails += check(
        "费用" not in joined and "token" not in joined.lower() and "提示词" not in joined,
        "选词文档无费用/token/提示词",
        "选词文档含禁词",
    )

    log("--- 2 普通版：截取负担说明；去掉 [360字左右] ---")
    data2, warn2 = build_kind_docx("plain", topic, SAMPLE_FULL)
    fails += check(warn2 is None, "完整三段式无 warning", f"warn={warn2}")
    doc2 = Document(io.BytesIO(data2))
    title2 = doc2.paragraphs[0]
    fails += check(
        title2.text == "【普通版-负担说明】雅各的变化",
        f"普通版首段={title2.text!r}",
        f"普通版首段不符={title2.text!r}",
    )
    fails += check(
        title2.alignment == WD_ALIGN_PARAGRAPH.CENTER
        and bool(title2.runs)
        and title2.runs[0].bold
        and title2.runs[0].font.size == Pt(16),
        "普通版首段居中加粗16pt",
        "普通版首段样式不符",
    )
    body2 = _paras(doc2)[1:]
    fails += check(
        body2 == ["第一行负担正文。", "第二行负担正文。", "第三行负担正文。", ""],
        f"负担正文={body2!r}",
        f"负担正文不符={body2!r}",
    )
    full_text = "\n".join(_paras(doc2))
    fails += check(
        "【职事界定】" not in full_text
        and "【真理脉络】" not in full_text
        and "【负担说明】" not in full_text
        and "[360字左右]" not in full_text,
        "不含前两段标题、负担说明标题与格式提示",
        f"full={full_text!r}",
    )
    fails += check(
        "费用" not in full_text and "token" not in full_text.lower(),
        "普通版无费用/token",
        "普通版含禁词",
    )

    data_ref, _ = build_kind_docx("ref", topic, SAMPLE_FULL)
    doc_ref = Document(io.BytesIO(data_ref))
    fails += check(
        doc_ref.paragraphs[0].text == "【加词条定义版-负担说明】雅各的变化",
        f"加词条首段={doc_ref.paragraphs[0].text!r}",
        f"加词条首段不符={doc_ref.paragraphs[0].text!r}",
    )

    # 标题后无换行、直接跟提示
    inline = "【职事界定】a\n【真理脉络】b\n【负担说明】[360字左右]\n正文甲\n正文乙"
    body_i, w_i = extract_burden_export(inline)
    fails += check(
        w_i is None and body_i == "正文甲\n正文乙",
        f"去提示后 body={body_i!r}",
        f"去提示失败 body={body_i!r} warn={w_i}",
    )

    log("--- 3 无【负担说明】→ 全文 + warning ---")
    data3, warn3 = build_kind_docx("ref", topic, SAMPLE_NO_BURDEN)
    fails += check(
        warn3 == "未找到【负担说明】标题，已导出全文",
        f"warning={warn3!r}",
        f"warning 不符={warn3!r}",
    )
    doc3 = Document(io.BytesIO(data3))
    body3 = _paras(doc3)[1:]
    expected3 = SAMPLE_NO_BURDEN.replace("\r\n", "\n").split("\n")
    fails += check(body3 == expected3, f"全文导出 body={body3!r}", f"body={body3!r}")

    log("--- 4 文件名与 Content-Disposition ---")
    for kind, name in (
        ("select", "【选词结果】雅各的变化.docx"),
        ("plain", "【普通版-负担说明】雅各的变化.docx"),
        ("ref", "【加词条定义版-负担说明】雅各的变化.docx"),
    ):
        got = kind_docx_filename(kind, topic)
        fails += check(got == name, f"{kind} → {got}", f"{kind} got {got} expected {name}")

    dirty = "雅各/变化:试验"
    safe_name = kind_docx_filename("plain", dirty)
    fails += check(
        "/" not in safe_name
        and ":" not in safe_name
        and safe_name.startswith("【普通版-负担说明】")
        and safe_name.endswith(".docx"),
        f"非法字符替换 → {safe_name}",
        f"未替换 → {safe_name}",
    )
    long_topic = "甲" * 120
    long_name = kind_docx_filename("plain", long_topic)
    fails += check(
        long_name.startswith("【普通版-负担说明】")
        and long_name.endswith(".docx")
        and len(long_name) <= 100
        and "【普通版-负担说明】" in long_name,
        f"过长主题截断保前缀 → len={len(long_name)} name={long_name}",
        f"截断异常 → {long_name!r} len={len(long_name)}",
    )
    cd = content_disposition(kind_docx_filename("select", topic))
    fails += check(
        "filename*=UTF-8''" in cd,
        f"Content-Disposition 含 filename*: {cd}",
        f"缺 filename*: {cd}",
    )
    encoded = cd.split("filename*=UTF-8''", 1)[1]
    decoded = unquote(encoded)
    fails += check(
        decoded == "【选词结果】雅各的变化.docx",
        f"解码文件名={decoded}",
        f"解码失败={decoded!r}",
    )
    # 选词结果文件名与首段未变
    fails += check(
        kind_docx_filename("select", topic) == "【选词结果】雅各的变化.docx"
        and Document(io.BytesIO(data)).paragraphs[0].text == topic,
        "选词结果文件名与首段内容未变",
        "选词结果被改动",
    )

    log("--- 5 现有整体导出仍可用（冒烟，格式未改断言结构）---")
    old = build_sc_docx(
        {
            "topic": topic,
            "final_terms": [{"name": "生命", "tier": "最内在", "reason": "x"}],
            "skipped": [],
            "merged": [],
            "plain_result": "普通全文",
            "ref_result": "加词条全文",
        }
    )
    old_doc = Document(io.BytesIO(old))
    old_texts = _paras(old_doc)
    fails += check(
        any("费用" in t for t in old_texts) and any("篇题" in t for t in old_texts),
        "整体导出仍含费用行与篇题（未改坏）",
        f"整体导出异常 {old_texts[:5]}",
    )

    log("--- 6 前端构建由外层执行；检查按钮文案 ---")
    vue = (_REPO / "front_mic" / "frontend" / "src" / "components" / "toolbox" / "PanAI4SC.vue").read_text(
        encoding="utf-8"
    )
    fails += check(
        vue.count("导出普通版-负担说明 docx") >= 2
        and vue.count("导出加词条定义版-负担说明 docx") >= 2
        and "导出选词结果 docx" in vue,
        "三处按钮文案已更新（生成区+详情各一对 + 选词）",
        "按钮文案未齐",
    )
    log("(npm run build)")

    log("")
    log(f"=== RESULT fails={fails} ===")
    _LOG.parent.mkdir(parents=True, exist_ok=True)
    # 追加不覆盖
    with _LOG.open("a", encoding="utf-8") as f:
        f.write("\n")
        f.write("\n".join(_lines) + "\n")
    log(f"appended {_LOG}")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
