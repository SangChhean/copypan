# -*- coding: utf-8 -*-
"""对照实验：B 摘录版 / C 直写版 vs run #19（参考优先版 v2）。

用法（在 back_mic/backend）:
    python -m panai4.sc.exp_bc            # dry-run：只组装提示词
    python -m panai4.sc.exp_bc --run      # 真实 API（由用户执行）
    python -m panai4.sc.exp_bc --run-id 19
"""
from __future__ import annotations

import argparse
import asyncio
import difflib
import json
import re
import sys
from pathlib import Path

_BACKEND = Path(__file__).resolve().parents[2]
_REPO = _BACKEND.parents[1]
_OUT_DIR = _REPO / "docs" / "panai4_sc" / "exp_bc"
_LOG = _REPO / "logs" / "sc_exp_bc.log"
_V1_PATH = Path(__file__).resolve().parent / "prompts" / "ref_prompt_v1_block_only.txt"

_lines: list[str] = []


def log(msg: str = "") -> None:
    _lines.append(msg)
    try:
        print(msg)
    except UnicodeEncodeError:
        sys.stdout.buffer.write((msg + "\n").encode("utf-8", errors="replace"))


def flush_log() -> None:
    _LOG.parent.mkdir(parents=True, exist_ok=True)
    _LOG.write_text("\n".join(_lines) + "\n", encoding="utf-8")


class ReplaceError(RuntimeError):
    pass


def exact_replace_once(text: str, old: str, new: str, *, label: str) -> str:
    n = text.count(old)
    if n != 1:
        raise ReplaceError(f"{label}：在底稿中出现 {n} 次（期望 1），停止")
    return text.replace(old, new, 1)


def exact_delete_once(text: str, old: str, *, label: str) -> str:
    return exact_replace_once(text, old, "", label=label)


# ---------- B / C 替换文案 ----------

B2_NEW = """第二步：摘录参考材料——
·逐条读文首【词条定义】，找出与本题直接相关的职事说法，逐字摘录定义中的原句。每条摘录必须是该词条定义里连续出现的原文，不改字、不缩写、不合并两处；一条只摘一句或一小段。
·每条格式：「参考材料·词条名｜原句：「……」」。同一词条可摘多条；摘录条数不限，以够写出本题的真理脉络为准；与本题无关的词条不摘。
·摘录之后，参考材料里没有、但本题确需的说法，才可用你所知道的职事信息补列，标「补」，并注明依据（经文，或职事材料的出处范围，至少到书卷或主题一级，不可编造篇名）与记忆程度（原用语／大意／印象；印象须写得出大约出自哪一类职事材料）。补列条数不得超过摘录条数的三分之一。
·须是职事信息方向的说法，不用一般基督教的解释补充；写成本题的具体形式，换一个篇题也能照搬的通用说法不列。
"""

B3_OLD = "依据类型分三种：经文明文、职事明言、自行推论。"
B3_NEW = (
    "依据类型分四种：经文明文、职事明言、参考材料、自行推论。"
    "「参考材料」指第二步的摘录，依据写「参考材料：词条名」，可作主干环。"
)

B4_NEW = """1. 参考材料·词条名｜原句：「……」

2. 参考材料·词条名｜原句：「……」

补1. ……（依据：……｜原用语／大意／印象）
"""

B5_OLD = "（经文明文／职事明言：出处）"
B5_NEW = "（经文明文／职事明言／参考材料：词条名）"

C0_OLD = (
    "· 凡依据这些定义写出的界定，在【职事界定】的「依据」处写「参考材料·词条名」；"
    "取自定义原文的说法，记忆程度可标「原用语」。"
)
C0_NEW = "· 用到的词条，在文末【用到的参考材料】里列出词条名。"


def _scan_angle_block(draft: str) -> str:
    """B1：含扫描角度上方分隔线，到【写作步骤】上方分隔线之前。"""
    start = draft.find("━━━━━━━━━━━━━━━━━━━\n🔍扫描角度（供第二步第(3)轮查漏使用）")
    if start < 0:
        raise ReplaceError("B1：找不到扫描角度标题块")
    end = draft.find("━━━━━━━━━━━━━━━━━━━\n【写作步骤】", start)
    if end < 0:
        raise ReplaceError("B1：找不到【写作步骤】标题块")
    return draft[start:end]


def _step2_block(draft: str) -> str:
    start = draft.find("第二步：列出职事界定——分三轮进行：")
    if start < 0:
        raise ReplaceError("B2：找不到第二步标题")
    end = draft.find("第三步：定核心命题", start)
    if end < 0:
        raise ReplaceError("B2：找不到第三步标题")
    return draft[start:end]


def _b4_block(draft: str) -> str:
    start = draft.find("1. ……（依据：……｜原用语／大意／印象｜角度X）")
    if start < 0:
        raise ReplaceError("B4：找不到输出格式职事界定示例行")
    end_marker = "无内容的角度：……"
    end = draft.find(end_marker, start)
    if end < 0:
        raise ReplaceError("B4：找不到「无内容的角度：……」")
    return draft[start : end + len(end_marker)]


def apply_B(draft: str) -> tuple[str, list[str]]:
    applied: list[str] = []
    text = draft
    block = _scan_angle_block(text)
    text = exact_delete_once(text, block, label="B1")
    applied.append("B1")
    s2 = _step2_block(text)
    text = exact_replace_once(text, s2, B2_NEW, label="B2")
    applied.append("B2")
    text = exact_replace_once(text, B3_OLD, B3_NEW, label="B3")
    applied.append("B3")
    b4 = _b4_block(text)
    text = exact_replace_once(text, b4, B4_NEW.rstrip("\n"), label="B4")
    applied.append("B4")
    text = exact_replace_once(text, B5_OLD, B5_NEW, label="B5")
    applied.append("B5")
    return text, applied


def apply_C(draft: str) -> tuple[str, list[str]]:
    """与页面默认直写版（v3）共用 prompt_v3.apply_ref_v3_transforms。"""
    from panai4.sc.prompt_v3 import apply_ref_v3_transforms

    return apply_ref_v3_transforms(draft)


def extract_term_defs_section(full: str) -> str:
    m = re.search(r"【词条定义】\n(.*?)\n━━━━━━━━━━━━━━━━━━━", full, re.S)
    if not m:
        raise RuntimeError("无法从文本解析【词条定义】段")
    return m.group(1)


def extract_ref_header(full: str) -> str:
    """参考块 + 词条定义（到词条定义后第一条分隔线为止，含该分隔线）。"""
    m = re.search(r"^(.*?【词条定义】\n.*?\n━━━━━━━━━━━━━━━━━━━\n)", full, re.S)
    if not m:
        raise RuntimeError("无法截取参考块+词条定义段")
    return m.group(1)


def load_run(run_id: int) -> dict:
    from panai4.sc import store
    from panai4 import db as panai4_db

    store.ensure_tables()
    row = panai4_db.connect().execute(
        "SELECT * FROM sc_runs WHERE id = ?", (run_id,)
    ).fetchone()
    if not row:
        raise RuntimeError(f"找不到 run #{run_id}")
    item = dict(row)
    for key in ("final_terms_json", "skipped_json", "merged_json"):
        raw = item.get(key)
        name = key.replace("_json", "")
        if raw:
            try:
                item[name] = json.loads(raw)
            except json.JSONDecodeError:
                item[name] = []
        else:
            item[name] = []
    return item


def build_term_defs_for_run(run: dict) -> str:
    from panai4.sc.data import build_term_defs, load_sc_data

    data = load_sc_data(force=True)
    terms = run.get("final_terms") or []
    names = [t["name"] if isinstance(t, dict) else t for t in terms]
    text, skipped, merged, *_ = build_term_defs(names, data)
    # 与存档 skipped/merged 仅作信息；TERM_DEFS 以重生成为准
    log(f"build_term_defs skipped={skipped} merged={merged}")
    log(f"archive skipped={run.get('skipped')} merged={run.get('merged')}")
    return text


def fill_prompt(template: str, topic: str, term_defs: str) -> str:
    from panai4.sc.data import fill_ref_prompt, load_sc_data

    data = load_sc_data(force=True)
    return fill_ref_prompt(template, topic, term_defs, data)


def remaining_braces(text: str) -> list[str]:
    return sorted(set(re.findall(r"\{([A-Za-z_]+)\}", text)))


def write_diff(label: str, a: str, b: str, a_name: str, b_name: str) -> None:
    log(f"\n===== difflib {label} =====")
    for line in difflib.unified_diff(
        a.splitlines(),
        b.splitlines(),
        fromfile=a_name,
        tofile=b_name,
        lineterm="",
    ):
        log(line)


def dry_run(run_id: int) -> int:
    log(f"=== SC EXP BC dry-run (run_id={run_id}) ===")
    _OUT_DIR.mkdir(parents=True, exist_ok=True)

    run = load_run(run_id)
    topic = run.get("topic") or ""
    archived_full = run.get("ref_prompt_full") or ""
    if not archived_full:
        raise RuntimeError(f"run #{run_id} 无 ref_prompt_full")

    term_defs = build_term_defs_for_run(run)
    arch_defs = extract_term_defs_section(archived_full)
    if term_defs.rstrip("\n") != arch_defs.rstrip("\n"):
        log("FAIL: 重生 TERM_DEFS 与存档【词条定义】不一致，停止")
        # 找首处差异
        a, b = term_defs.rstrip("\n"), arch_defs.rstrip("\n")
        for i, (ca, cb) in enumerate(zip(a, b)):
            if ca != cb:
                log(f"  first diff @{i}: new={ca!r} arch={cb!r} context_new={a[max(0,i-20):i+20]!r}")
                break
        else:
            log(f"  length new={len(a)} arch={len(b)}")
        flush_log()
        return 1
    log("PASS: TERM_DEFS 与 run 存档逐字一致")

    draft = _V1_PATH.read_text(encoding="utf-8")
    if "{TERM_DEFS}" not in draft or "{query}" not in draft:
        raise RuntimeError("v1 底稿缺少占位符")

    try:
        tpl_b, applied_b = apply_B(draft)
        tpl_c, applied_c = apply_C(draft)
    except ReplaceError as e:
        log(f"STOP: {e}")
        flush_log()
        return 1

    log(f"B 已应用替换: {', '.join(applied_b)}")
    log(f"C 已应用替换: {', '.join(applied_c)}")

    prompt_b = fill_prompt(tpl_b, topic, term_defs)
    prompt_c = fill_prompt(tpl_c, topic, term_defs)
    (_OUT_DIR / "prompt_B.txt").write_text(prompt_b, encoding="utf-8")
    (_OUT_DIR / "prompt_C.txt").write_text(prompt_c, encoding="utf-8")
    log(f"wrote {_OUT_DIR / 'prompt_B.txt'} chars={len(prompt_b)}")
    log(f"wrote {_OUT_DIR / 'prompt_C.txt'} chars={len(prompt_c)}")

    for name, text in (("B", prompt_b), ("C", prompt_c)):
        left = remaining_braces(text)
        bad = [x for x in left if x in {"DEF_WAI", "DEF_NEI", "DEF_ZUINEI", "TERM_DEFS", "query"}]
        if bad:
            log(f"FAIL: {name} 仍有占位符 {bad}")
            flush_log()
            return 1
        log(f"PASS: {name} 无关键占位符 leftover={left}")

    # 底稿→B/C 的模板级 diff（未填占位符，便于阅读）
    write_diff("draft→B", draft, tpl_b, "v1_draft", "tpl_B")
    write_diff("draft→C", draft, tpl_c, "v1_draft", "tpl_C")

    # 参考块+词条定义 vs #19
    arch_header = extract_ref_header(archived_full)
    b_header = extract_ref_header(prompt_b)
    c_header = extract_ref_header(prompt_c)
    if b_header == arch_header:
        log("PASS: prompt_B 参考块+词条定义 与 #19 逐字一致")
    else:
        log("FAIL: prompt_B 参考块+词条定义 与 #19 不一致")
        write_diff("arch↔B_header", arch_header, b_header, "arch", "B")

    # C：去掉 C0 那一行后再比
    arch_wo_c0 = arch_header.replace(C0_OLD, "", 1) if C0_OLD in arch_header else arch_header
    c_wo_c0 = c_header.replace(C0_NEW, "", 1) if C0_NEW in c_header else c_header
    # 规范化：C0 删除后可能留下双换行差异，只比「其余」——把 C0 行换成同一占位
    arch_norm = arch_header.replace(C0_OLD, "⟦C0⟧", 1)
    c_norm = c_header.replace(C0_NEW, "⟦C0⟧", 1)
    if arch_norm == c_norm:
        log("PASS: prompt_C 参考块+词条定义（除 C0 外）与 #19 一致")
    else:
        log("REPORT: prompt_C 参考块除 C0 外仍有差异")
        write_diff("arch↔C_header_norm", arch_norm, c_norm, "arch_norm", "C_norm")

    log("=== dry-run done ===")
    flush_log()
    return 0


# ---------- --run ----------

def _burden_body(text: str) -> str:
    from panai4.sc.export_sc import extract_burden_export

    body, _ = extract_burden_export(text or "")
    return body


def _burden_char_count(text: str) -> int:
    body = _burden_body(text)
    return len(re.sub(r"\s+", "", body))


def _parse_b_quotes(text: str) -> list[tuple[str, str]]:
    """返回 [(词条名, 原句), ...]。"""
    # 参考材料·词条名｜原句：「……」
    pat = re.compile(r"参考材料·([^｜|]+)[｜|]原句[：:]「([^」]*)」")
    return [(m.group(1).strip(), m.group(2)) for m in pat.finditer(text or "")]


def _norm_ws(s: str) -> str:
    return re.sub(r"[\s\u3000]+", "", s or "")


def _outline_ring_count(text: str) -> int | str:
    from panai4.sc.citation import split_sections

    sec = split_sections(text or "").get("【真理脉络】") or ""
    # 计「数字.」行，排除 0. 与应用可计入；用户要环数——计 1. 起的主干，含应用
    nos = []
    for line in sec.splitlines():
        m = re.match(r"^(\d+)\.", line.lstrip())
        if m:
            nos.append(int(m.group(1)))
    if not nos:
        return "无"
    # 环数通常不含第0行
    rings = [n for n in nos if n > 0]
    return len(rings) if rings else 0


def _parse_c_materials(text: str) -> list[str]:
    m = re.search(r"【用到的参考材料】\s*(.*?)(?:\n【|\Z)", text or "", re.S)
    if not m:
        return []
    chunk = m.group(1).strip()
    parts = re.split(r"[、，,]", chunk)
    return [p.strip().strip("「」『』") for p in parts if p.strip()]


def _export_burden_docx(topic: str, raw: str, prefix: str, path: Path) -> None:
    from docx import Document
    from panai4.export_docx import safe_filename_part
    from panai4.sc.export_sc import _add, _set_title_paragraph, extract_burden_export

    body, _ = extract_burden_export(raw)
    doc = Document()
    title = f"{prefix}{topic}"
    if doc.paragraphs:
        _set_title_paragraph(doc.paragraphs[0], title)
    else:
        p = doc.add_paragraph()
        _set_title_paragraph(p, title)
    for line in body.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        _add(doc, line)
    max_topic = max(1, 100 - len(prefix) - len(".docx"))
    fname = f"{prefix}{safe_filename_part(topic, max_topic)}.docx"
    out = path / fname
    buf_path = out
    doc.save(str(buf_path))
    log(f"wrote {buf_path}")


async def _call_both(prompt_b: str, prompt_c: str, topic: str):
    from panai4.sc.llm_call import sc_complete, usage_dict
    from panai4.sc.cost import calc_cost_usd, format_usd

    rb, rc = await asyncio.gather(
        sc_complete(prompt_b, topic=topic),
        sc_complete(prompt_c, topic=topic),
    )
    ub = usage_dict(rb)
    uc = usage_dict(rc)
    ub["cost_usd_calc"] = calc_cost_usd(ub)
    uc["cost_usd_calc"] = calc_cost_usd(uc)
    ub["cost_display"] = format_usd(ub["cost_usd_calc"])
    uc["cost_display"] = format_usd(uc["cost_usd_calc"])
    return rb, rc, ub, uc


def run_api(run_id: int) -> int:
    log(f"=== SC EXP BC --run (run_id={run_id}) ===")
    _OUT_DIR.mkdir(parents=True, exist_ok=True)

    run = load_run(run_id)
    topic = run.get("topic") or ""
    term_defs = build_term_defs_for_run(run)
    arch_defs = extract_term_defs_section(run.get("ref_prompt_full") or "")
    if term_defs.rstrip("\n") != arch_defs.rstrip("\n"):
        log("FAIL: TERM_DEFS 不一致，拒绝 --run")
        flush_log()
        return 1

    draft = _V1_PATH.read_text(encoding="utf-8")
    try:
        tpl_b, _ = apply_B(draft)
        tpl_c, _ = apply_C(draft)
    except ReplaceError as e:
        log(f"STOP: {e}")
        flush_log()
        return 1

    prompt_b = fill_prompt(tpl_b, topic, term_defs)
    prompt_c = fill_prompt(tpl_c, topic, term_defs)
    (_OUT_DIR / "prompt_B.txt").write_text(prompt_b, encoding="utf-8")
    (_OUT_DIR / "prompt_C.txt").write_text(prompt_c, encoding="utf-8")

    rb, rc, ub, uc = asyncio.run(_call_both(prompt_b, prompt_c, topic))
    (_OUT_DIR / "out_B.txt").write_text(rb.text or "", encoding="utf-8")
    (_OUT_DIR / "out_C.txt").write_text(rc.text or "", encoding="utf-8")
    (_OUT_DIR / "usage_B.json").write_text(
        json.dumps(ub, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (_OUT_DIR / "usage_C.json").write_text(
        json.dumps(uc, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    log("wrote out_B/out_C + usage_*.json")

    _write_report(run, rb.text or "", rc.text or "", ub, uc)
    _export_burden_docx(topic, rb.text or "", "【B摘录版-负担说明】", _OUT_DIR)
    _export_burden_docx(topic, rc.text or "", "【C直写版-负担说明】", _OUT_DIR)

    flush_log()
    return 0


def _write_report(run: dict, out_b: str, out_c: str, ub: dict, uc: dict) -> None:
    from panai4.sc.citation import compute_citation_stats
    from panai4.sc.cost import format_usd
    from panai4.sc.data import load_sc_data

    data = load_sc_data(force=True)
    terms = run.get("final_terms") or []
    skipped = run.get("skipped") or []
    merged = run.get("merged") or []
    out_19 = run.get("ref_result") or ""
    usage_19 = {}
    if run.get("ref_usage_json"):
        try:
            usage_19 = json.loads(run["ref_usage_json"]) if isinstance(run["ref_usage_json"], str) else run["ref_usage_json"]
        except json.JSONDecodeError:
            usage_19 = {}

    # B quote check
    quotes = _parse_b_quotes(out_b)
    name_to_def: dict[str, str] = {}
    for t in terms:
        name = t["name"] if isinstance(t, dict) else t
        c = data.content_3a.get(name)
        if c:
            name_to_def[name] = str(c)
    match_n = 0
    mismatch = []
    for name, quote in quotes:
        src = name_to_def.get(name) or ""
        if _norm_ws(quote) and _norm_ws(quote) in _norm_ws(src):
            match_n += 1
        else:
            mismatch.append((name, quote[:40]))

    cite_b = compute_citation_stats(
        ref_output=out_b,
        final_terms=terms,
        skipped=skipped,
        merged=merged,
        content_3a=data.content_3a,
        tier_of=data.tier_of,
    )
    cite_19 = compute_citation_stats(
        ref_output=out_19,
        final_terms=terms,
        skipped=skipped,
        merged=merged,
        content_3a=data.content_3a,
        tier_of=data.tier_of,
    )

    mats_c = _parse_c_materials(out_c)
    final_names = {t["name"] if isinstance(t, dict) else t for t in terms}
    unknown_c = [n for n in mats_c if n not in final_names]
    by_tier_c = {"最内在": 0, "内在": 0, "外在": 0}
    for n in mats_c:
        if n in final_names:
            tier = data.tier_of.get(n) or "外在"
            if tier in by_tier_c:
                by_tier_c[tier] += 1

    def cited_count(stats: dict) -> int:
        bt = stats.get("by_tier") or {}
        return sum(bt.get(t, {}).get("in_outline", 0) + bt.get(t, {}).get("defs_only", 0) for t in ("最内在", "内在", "外在"))

    def tier_cite(stats: dict) -> str:
        bt = stats.get("by_tier") or {}
        parts = []
        for t in ("最内在", "内在", "外在"):
            b = bt.get(t) or {}
            parts.append(
                f"{t} 入{b.get('in_outline',0)}/仅{b.get('defs_only',0)}/未{b.get('uncited',0)}（共{b.get('total',0)}）"
            )
        return "；".join(parts)

    cost_19 = run.get("cost_ref_usd")
    lines = [
        "# SC 对照实验报告：#19(v2) / B摘录版 / C直写版",
        "",
        f"篇题：{run.get('topic')}",
        "",
        "## 1. 负担说明字数",
        f"- #19：{_burden_char_count(out_19)}",
        f"- B：{_burden_char_count(out_b)}",
        f"- C：{_burden_char_count(out_c)}",
        "",
        "## 2. B 原句子串匹配",
        f"- 总条数：{len(quotes)}；完全匹配：{match_n}；不匹配：{len(mismatch)}",
    ]
    for name, head in mismatch:
        lines.append(f"  - {name}｜{head}")
    lines += [
        "",
        "## 3. B 引用统计",
        tier_cite(cite_b),
        f"note={cite_b.get('note')}",
        "",
        "## 4. C 用到的参考材料",
        f"列表：{'、'.join(mats_c) if mats_c else '（无/未解析）'}",
        f"分档计数：{by_tier_c}",
        f"不在最终名单：{unknown_c or '无'}",
        "",
        "## 5. 对照表",
        "",
        "| 项目 | #19 v2 | B 摘录版 | C 直写版 |",
        "|---|---|---|---|",
        f"| 费用 | {format_usd(float(cost_19) if cost_19 is not None else None)} | {ub.get('cost_display')} | {uc.get('cost_display')} |",
        f"| 输入 token | {usage_19.get('input_tokens', '—')} | {ub.get('input_tokens')} | {uc.get('input_tokens')} |",
        f"| 输出 token | {usage_19.get('output_tokens', '—')} | {ub.get('output_tokens')} | {uc.get('output_tokens')} |",
        f"| 负担说明字数 | {_burden_char_count(out_19)} | {_burden_char_count(out_b)} | {_burden_char_count(out_c)} |",
        f"| 被引用词条数 | {cited_count(cite_19)} | {cited_count(cite_b)} | {len([n for n in mats_c if n in final_names])} |",
        f"| 真理脉络环数 | {_outline_ring_count(out_19)} | {_outline_ring_count(out_b)} | 无 |",
        "",
        "### 引用分档",
        f"- #19：{tier_cite(cite_19)}",
        f"- B：{tier_cite(cite_b)}",
        f"- C 列出分档：{by_tier_c}",
        "",
        "## 6. 三份负担说明正文",
        "",
        "### #19 参考优先版（v2）",
        _burden_body(out_19),
        "",
        "### B 摘录版",
        _burden_body(out_b),
        "",
        "### C 直写版",
        _burden_body(out_c),
        "",
    ]
    path = _OUT_DIR / "report.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    log(f"wrote {path}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="SC B/C 对照实验")
    parser.add_argument("--run", action="store_true", help="真实调用 API（默认 dry-run）")
    parser.add_argument("--run-id", type=int, default=19)
    args = parser.parse_args(argv)
    try:
        if args.run:
            return run_api(args.run_id)
        return dry_run(args.run_id)
    except Exception as e:
        log(f"ERROR: {e}")
        flush_log()
        raise


if __name__ == "__main__":
    sys.exit(main())
