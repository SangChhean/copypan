# -*- coding: utf-8 -*-
"""SC 测试第 1 段：从三档名单 docx 生成分类 JSON，并对齐 3a / Neo4j。

只读输入；产出写到 docs/panai4_graph/词条三档分类.json、
docs/panai4_sc/stage1_report.md、docs/panai4_sc/null_terms.csv、
logs/sc_stage1.log。

用法（在 back_mic/backend 下）:
    python -m panai4.sc.stage1
"""
from __future__ import annotations

import csv
import json
import os
import random
import re
import statistics
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from neo4j import READ_ACCESS, GraphDatabase

_BACKEND = Path(__file__).resolve().parents[2]
_REPO = _BACKEND.parents[1]
_DOCX = _REPO / "docs" / "2373词条_外在内在最内在分类（10.1）.docx"
_JSON_3A = _REPO / "docs" / "3a_export.json"
_OUT_TIERS = _REPO / "docs" / "panai4_graph" / "词条三档分类.json"
_OUT_DIR = _REPO / "docs" / "panai4_sc"
_OUT_REPORT = _OUT_DIR / "stage1_report.md"
_OUT_NULL_CSV = _OUT_DIR / "null_terms.csv"
_OUT_LOG = _REPO / "logs" / "sc_stage1.log"

_PAREN_RE = re.compile(r"^([^（）()]+)（([^（）()]+)）$")
_ANY_PAREN_RE = re.compile(r"[（）()]")
_TIER_ORDER = ("外在", "内在", "最内在")
_HEADER_MARKERS = frozenset({"分类", "词条", "名称", "档", "数量", "占比"})
_SEED = 20261002
_SAMPLE_N = 30


class AmbiguousTermError(RuntimeError):
    pass


def _log_lines(lines: list[str], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _row_cells(row) -> list[str]:
    """同行去合并单元格重复后的单元格文本（已 strip）。"""
    cells: list[str] = []
    prev: str | None = None
    for cell in row.cells:
        text = cell.text.strip()
        if prev is not None and text == prev:
            continue
        prev = text
        cells.append(text)
    return cells


def _is_header_row(cells: list[str]) -> bool:
    nonempty = [c for c in cells if c]
    if not nonempty:
        return False
    if nonempty[0] in _HEADER_MARKERS:
        return True
    if all(c in _HEADER_MARKERS for c in nonempty):
        return True
    return False


def _pull_name_table(table, label: str) -> list[str]:
    """从词条表抽取名字：跳过表头行；一词一格。"""
    rows = [_row_cells(row) for row in table.rows]
    start = 1 if rows and _is_header_row(rows[0]) else 0
    terms: list[str] = []
    for cells in rows[start:]:
        for name in cells:
            if name == "":
                continue
            parts = [ln.strip() for ln in name.replace("\r", "").split("\n") if ln.strip()]
            if len(parts) != 1:
                raise AmbiguousTermError(
                    f"「{label}」单元格含多个名字或空行混杂: {parts!r}"
                )
            name = parts[0]
            if re.search(r"^\d+[\.、．)]", name) or re.search(r"^（\d+）", name):
                raise AmbiguousTermError(f"「{label}」疑似编号: {name!r}")
            if "、" in name or "," in name or "；" in name or ";" in name:
                raise AmbiguousTermError(f"「{label}」疑似多名并列: {name!r}")
            terms.append(name)
    return terms


def extract_docx_tiers(docx_path: Path) -> tuple[dict[str, list[str]], dict]:
    """返回 {档: [原词条…]}（未拆括号），以及结构说明 dict。

    新版 docx（10.1）布局：
    - 表0：统计（不取用）
    - 最内在 → 表1；内在 → 表2
    - 外在 = 纯外在（表3）+ 由内在彰显出的外在（表4），顺序拼接
    """
    doc = Document(str(docx_path))
    if len(doc.tables) != 5:
        raise AmbiguousTermError(f"期望 5 个表格（统计+最内在+内在+一般外在+外在的内在），实际 {len(doc.tables)}")

    headings = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading")]
    extra_paras = [
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip() and not (p.style.name.startswith("Heading") or p.style.name == "Title")
    ]
    title_paras = [p.text.strip() for p in doc.paragraphs if p.style.name == "Title" and p.text.strip()]

    most = _pull_name_table(doc.tables[1], "最内在")
    inner = _pull_name_table(doc.tables[2], "内在")
    outer_general = _pull_name_table(doc.tables[3], "外在/纯外在")
    outer_inner = _pull_name_table(doc.tables[4], "外在/由内在彰显出的外在")
    outer = outer_general + outer_inner

    tiers: dict[str, list[str]] = {"外在": outer, "内在": inner, "最内在": most}
    meta = {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "headings": headings,
        "extra_paragraphs": title_paras + extra_paras,
        "raw_counts": {k: len(v) for k, v in tiers.items()},
        "paren_terms": {},
        "empty_cells": {},
        "outer_parts": {"纯外在": len(outer_general), "由内在彰显出的外在": len(outer_inner)},
        "stats_table_skipped": True,
        "name_list_patch": "未配置（本脚本无名单修正逻辑）",
    }
    for tier, terms in tiers.items():
        meta["paren_terms"][tier] = [t for t in terms if _ANY_PAREN_RE.search(t)]

    return tiers, meta


def split_paren(name: str) -> list[str]:
    if not _ANY_PAREN_RE.search(name):
        return [name]
    m = _PAREN_RE.match(name)
    if not m:
        raise AmbiguousTermError(f"无法按单层全角括号拆开: {name!r}")
    outer, inner = m.group(1).strip(), m.group(2).strip()
    if not outer or not inner:
        raise AmbiguousTermError(f"括号内外为空: {name!r}")
    if _ANY_PAREN_RE.search(outer) or _ANY_PAREN_RE.search(inner):
        raise AmbiguousTermError(f"拆完后仍含括号: {name!r} → {outer!r}/{inner!r}")
    return [outer, inner]


def expand_tiers(raw_tiers: dict[str, list[str]]) -> dict[str, list[str]]:
    expanded: dict[str, list[str]] = {k: [] for k in _TIER_ORDER}
    for tier in _TIER_ORDER:
        for name in raw_tiers[tier]:
            expanded[tier].extend(split_paren(name))
    return expanded


def load_3a(path: Path) -> tuple[dict[str, str | None], dict]:
    rows = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(rows, list):
        raise RuntimeError("3a_export.json 顶层不是列表")
    field_stats = {}
    if rows:
        for field in rows[0].keys():
            kinds = Counter()
            nonempty = 0
            for row in rows:
                if field not in row:
                    kinds["missing"] += 1
                    continue
                v = row[field]
                if v is None:
                    kinds["null"] += 1
                elif isinstance(v, str):
                    if v == "":
                        kinds["empty_str"] += 1
                    elif v.strip() == "":
                        kinds["whitespace"] += 1
                    else:
                        nonempty += 1
                        kinds["nonempty_str"] += 1
                else:
                    nonempty += 1
                    kinds[f"other:{type(v).__name__}"] += 1
            field_stats[field] = {"nonempty": nonempty, "kinds": dict(kinds)}

    by_title: dict[str, str | None] = {}
    for row in rows:
        title = str(row.get("title", "")).strip()
        raw = row.get("content_3a")
        if raw is None:
            content = None
        elif isinstance(raw, str) and raw.strip() == "":
            content = None
        else:
            content = raw if isinstance(raw, str) else str(raw)
        if title in by_title:
            raise AmbiguousTermError(f"3a title 重复: {title}")
        by_title[title] = content
    meta = {
        "len": len(rows),
        "fields": list(rows[0].keys()) if rows else [],
        "field_stats": field_stats,
        "title_count": len(by_title),
    }
    return by_title, meta


def read_neo4j_truth_points() -> set[str]:
    env_path = _BACKEND / ".env"
    if env_path.is_file():
        load_dotenv(env_path)
    uri = (os.environ.get("NEO4J_V4_URI") or "").strip()
    user = (os.environ.get("NEO4J_V4_USER") or "").strip()
    password = os.environ.get("NEO4J_V4_PASSWORD") or ""
    if not uri or not user:
        raise SystemExit("缺少 NEO4J_V4_URI / NEO4J_V4_USER")
    driver = GraphDatabase.driver(uri, auth=(user, password))
    try:
        driver.verify_connectivity()
        with driver.session(default_access_mode=READ_ACCESS) as session:
            rows = session.run("MATCH (n:Truth_Point) RETURN n.name AS name")
            return {r["name"] for r in rows}
    finally:
        driver.close()


def self_check(
    raw_tiers: dict[str, list[str]],
    expanded: dict[str, list[str]],
    names_3a: set[str],
    names_neo: set[str],
) -> tuple[list[tuple[str, bool, str]], list[str]]:
    """返回 [(项名, 通过, 详情), ...] 与日志行。"""
    results: list[tuple[str, bool, str]] = []
    log: list[str] = []

    sets = {t: set(expanded[t]) for t in _TIER_ORDER}
    # 1. pairwise empty intersection
    ok1 = True
    details1 = []
    for i, a in enumerate(_TIER_ORDER):
        for b in _TIER_ORDER[i + 1 :]:
            inter = sorted(sets[a] & sets[b])
            details1.append(f"{a}∩{b}={len(inter)}")
            if inter:
                ok1 = False
                details1.append(f"  sample={inter[:10]}")
    results.append(("1_pairwise_intersection_empty", ok1, "; ".join(details1)))
    log.append(f"[1] pairwise_intersection_empty={ok1} {'; '.join(details1)}")

    # also fail hard if same name in two tiers during build
    name_to_tiers: dict[str, list[str]] = {}
    for t in _TIER_ORDER:
        for n in expanded[t]:
            name_to_tiers.setdefault(n, []).append(t)
    multi = {n: ts for n, ts in name_to_tiers.items() if len(ts) > 1}
    if multi:
        raise AmbiguousTermError(f"同一名字出现在多档: {dict(list(multi.items())[:20])}")

    union = set().union(*(sets[t] for t in _TIER_ORDER))
    # 2. union vs 3a and neo4j
    only_tiers_3a = sorted(union - names_3a)
    only_3a = sorted(names_3a - union)
    only_tiers_neo = sorted(union - names_neo)
    only_neo = sorted(names_neo - union)
    ok2 = (
        len(union) == 2381
        and not only_tiers_3a
        and not only_3a
        and not only_tiers_neo
        and not only_neo
    )
    detail2 = (
        f"union={len(union)} 3a={len(names_3a)} neo={len(names_neo)}; "
        f"tiers-3a={len(only_tiers_3a)} 3a-tiers={len(only_3a)}; "
        f"tiers-neo={len(only_tiers_neo)} neo-tiers={len(only_neo)}"
    )
    results.append(("2_union_vs_3a_and_neo4j", ok2, detail2))
    log.append(f"[2] {detail2}")
    if only_tiers_3a:
        log.append(f"[2] only_tiers_not_3a sample={only_tiers_3a[:20]}")
    if only_3a:
        log.append(f"[2] only_3a_not_tiers sample={only_3a[:20]}")
    if only_tiers_neo:
        log.append(f"[2] only_tiers_not_neo sample={only_tiers_neo[:20]}")
    if only_neo:
        log.append(f"[2] only_neo_not_tiers sample={only_neo[:20]}")

    # 3. counts = raw + paren_pairs per tier; total +8
    ok3 = True
    detail3_parts = []
    total_pairs = 0
    for t in _TIER_ORDER:
        raw_n = len(raw_tiers[t])
        pairs = sum(1 for n in raw_tiers[t] if _ANY_PAREN_RE.search(n))
        total_pairs += pairs
        exp_n = len(expanded[t])
        expect = raw_n + pairs
        good = exp_n == expect
        ok3 = ok3 and good
        detail3_parts.append(f"{t}: raw={raw_n} pairs={pairs} expanded={exp_n} expect={expect} ok={good}")
    ok3 = ok3 and total_pairs == 8
    detail3 = "; ".join(detail3_parts) + f"; total_pairs={total_pairs}"
    results.append(("3_counts_raw_plus_paren_pairs", ok3, detail3))
    log.append(f"[3] {detail3}")

    # 4. each paren pair in same tier
    ok4 = True
    pair_details = []
    for t in _TIER_ORDER:
        for name in raw_tiers[t]:
            if not _ANY_PAREN_RE.search(name):
                continue
            outer, inner = split_paren(name)
            same = outer in sets[t] and inner in sets[t]
            other = [x for x in _TIER_ORDER if x != t and (outer in sets[x] or inner in sets[x])]
            ok4 = ok4 and same and not other
            pair_details.append(f"{name}@{t} same={same} other={other}")
    results.append(("4_paren_pairs_same_tier", ok4, "; ".join(pair_details)))
    log.append(f"[4] paren_pairs_same_tier={ok4}")
    for line in pair_details:
        log.append(f"[4] {line}")

    # 5. no parens in expanded names
    with_paren = sorted(n for n in union if _ANY_PAREN_RE.search(n))
    ok5 = not with_paren
    results.append(("5_no_parens_in_names", ok5, f"with_paren={len(with_paren)} sample={with_paren[:10]}"))
    log.append(f"[5] no_parens={ok5} count={len(with_paren)}")

    # 6. raw total 2373, has 圣经中的奥秘, no 元首权柄
    raw_total = sum(len(raw_tiers[t]) for t in _TIER_ORDER)
    has_mystery = any("圣经中的奥秘" in raw_tiers[t] for t in _TIER_ORDER)
    has_headship = any("元首权柄" in raw_tiers[t] for t in _TIER_ORDER)
    where = [t for t in _TIER_ORDER if "圣经中的奥秘" in raw_tiers[t]]
    ok6 = raw_total == 2373 and has_mystery and not has_headship
    results.append(
        (
            "6_raw_total_2373_mystery_no_元首权柄",
            ok6,
            f"raw_total={raw_total} has_圣经中的奥秘={has_mystery} tiers={where} has_元首权柄={has_headship}",
        )
    )
    log.append(
        f"[6] raw_total={raw_total} has_圣经中的奥秘={has_mystery} tiers={where} "
        f"has_元首权柄={has_headship}"
    )

    return results, log


def write_stats(
    expanded: dict[str, list[str]],
    content_by_title: dict[str, str | None],
    checks: list[tuple[str, bool, str]],
    docx_meta: dict,
    meta_3a: dict,
) -> None:
    _OUT_DIR.mkdir(parents=True, exist_ok=True)

    null_rows: list[tuple[str, str]] = []
    lines: list[str] = []
    add = lines.append
    add("# SC 测试第 1 段报告（数据层）")
    add("")
    add("由 `python -m panai4.sc.stage1` 生成。不含 3a 定义原文。")
    add("")
    add("## 输入结构摘要")
    add("")
    add(f"- docx 标题：{docx_meta['headings']}")
    add(f"- docx 名单外段落（Title/说明，未取用）：{docx_meta['extra_paragraphs']}")
    add(f"- docx 各档原词条数：{docx_meta['raw_counts']}")
    if docx_meta.get("outer_parts"):
        add(f"- 外在子表：{docx_meta['outer_parts']}（按 docx 顺序拼入「外在」）")
    add(f"- docx 各档括号词条：{ {k: len(v) for k,v in docx_meta['paren_terms'].items()} }")
    add(f"- 名单修正：{docx_meta.get('name_list_patch', '无')}")
    add(f"- 3a 条数：{meta_3a['len']}；字段：{meta_3a['fields']}")
    add(f"- 3a `content_3a` 非空/空：见 field_stats")
    add("")
    add("## 自检")
    add("")
    add("| 项 | 通过 | 详情 |")
    add("|---|---|---|")
    for name, ok, detail in checks:
        add(f"| {name} | {'通过' if ok else '不通过'} | {detail} |")
    add("")
    add("## 各档与 3a 覆盖")
    add("")
    add("| 档 | 词条数 | 有 3a | NULL | NULL 占比 |")
    add("|---|---:|---:|---:|---:|")

    for tier in _TIER_ORDER:
        names = expanded[tier]
        has = 0
        null_n = 0
        for n in names:
            c = content_by_title.get(n)
            if c is None:
                null_n += 1
                null_rows.append((n, tier))
            else:
                has += 1
        total = len(names)
        ratio = (null_n / total) if total else 0.0
        add(f"| {tier} | {total} | {has} | {null_n} | {ratio:.2%} |")

    add("")
    add(f"- NULL 词条总数：**{len(null_rows)}**（明细见 `null_terms.csv`）")
    add("")

    # length stats per tier (only non-null)
    add("## 各档 3a 定义字数（仅有定义的词）")
    add("")
    add("| 档 | 样本数 | 最小 | 中位数 | 平均 | 最大 | 总和 |")
    add("|---|---:|---:|---:|---:|---:|---:|")
    for tier in _TIER_ORDER:
        lengths = []
        pairs = []
        for n in expanded[tier]:
            c = content_by_title.get(n)
            if c is None:
                continue
            L = len(c)
            lengths.append(L)
            pairs.append((n, L))
        if not lengths:
            add(f"| {tier} | 0 | - | - | - | - | 0 |")
            continue
        add(
            f"| {tier} | {len(lengths)} | {min(lengths)} | {statistics.median(lengths):.1f} | "
            f"{statistics.mean(lengths):.1f} | {max(lengths)} | {sum(lengths)} |"
        )
        top10 = sorted(pairs, key=lambda x: (-x[1], x[0]))[:10]
        add("")
        add(f"### {tier} 字数最多的前 10")
        add("")
        add("| 名字 | 字数 |")
        add("|---|---:|")
        for n, L in top10:
            add(f"| {n} | {L} |")
        add("")

    # random sample
    defined = [(n, t) for t in _TIER_ORDER for n in expanded[t] if content_by_title.get(n) is not None]
    rng = random.Random(_SEED)
    sample = rng.sample(defined, min(_SAMPLE_N, len(defined)))
    sample_sorted = sorted(sample, key=lambda x: (x[1], x[0]))
    total_chars = sum(len(content_by_title[n] or "") for n, _t in sample)
    add("## 随机样本（有定义，种子 20261002，n=30）")
    add("")
    add(f"- 抽取名单（名字 / 档）：")
    for n, t in sample_sorted:
        add(f"  - {n} / {t}")
    add(f"- 这 30 个词定义总字数：**{total_chars}**")
    add("")

    _OUT_REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")

    with _OUT_NULL_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["名字", "档"])
        for n, t in null_rows:
            w.writerow([n, t])


def main() -> int:
    raw_tiers, docx_meta = extract_docx_tiers(_DOCX)
    expanded = expand_tiers(raw_tiers)
    content_by_title, meta_3a = load_3a(_JSON_3A)
    names_neo = read_neo4j_truth_points()

    # write JSON (tier order keys as specified)
    payload = {t: expanded[t] for t in _TIER_ORDER}
    _OUT_TIERS.parent.mkdir(parents=True, exist_ok=True)
    _OUT_TIERS.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    checks, log = self_check(raw_tiers, expanded, set(content_by_title), names_neo)
    # prepend structure notes to log
    preamble = [
        "=== SC stage1 ===",
        f"docx={_DOCX}",
        f"3a={_JSON_3A}",
        f"out_tiers={_OUT_TIERS}",
        f"docx_raw_counts={docx_meta['raw_counts']}",
        f"docx_paren={ {k:v for k,v in docx_meta['paren_terms'].items()} }",
        f"expanded_counts={{ {', '.join(f'{t}:{len(expanded[t])}' for t in _TIER_ORDER)} }}",
        f"3a_len={meta_3a['len']} content_3a_kinds={meta_3a['field_stats'].get('content_3a')}",
        f"neo4j_Truth_Point={len(names_neo)}",
        "",
    ]
    _log_lines(preamble + log + [f"check_{n}={'PASS' if ok else 'FAIL'} {d}" for n, ok, d in checks], _OUT_LOG)

    write_stats(expanded, content_by_title, checks, docx_meta, meta_3a)

    failed = [n for n, ok, _ in checks if not ok]
    print("wrote", _OUT_TIERS)
    print("wrote", _OUT_REPORT)
    print("wrote", _OUT_NULL_CSV)
    print("wrote", _OUT_LOG)
    for n, ok, d in checks:
        print(f"{'PASS' if ok else 'FAIL'} {n}: {d}")
    if failed:
        return 1
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except AmbiguousTermError as exc:
        print("STOPPED:", exc, file=sys.stderr)
        raise SystemExit(2) from exc
