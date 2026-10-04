# -*- coding: utf-8 -*-
"""SC 测试页导出 docx。"""
from __future__ import annotations

import io
import json
import re
from typing import Literal
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from panai4.export_docx import safe_filename_part
from panai4.pipeline import BURDEN_MARKER, BURDEN_TAG

from .citation import STATUS_LABEL, TIER_ORDER, compute_citation_stats
from .cost import format_usd

_FONT = "Microsoft YaHei"

ExportKind = Literal["select", "plain", "ref"]

_KIND_LABEL = {
    "select": "选词结果",
    "plain": "普通版-负担说明",
    "ref": "加词条定义版-负担说明",
}

# 【负担说明】之后紧跟的格式提示，如 [360字左右]
_HINT_AFTER_BURDEN = re.compile(r"^[ \t]*\[[^\]]{0,30}\][ \t]*")


def _set_run_font(run, size_pt: float, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = _FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), _FONT)


def _add(
    doc: Document,
    text: str,
    *,
    bold: bool = False,
    size: float = 11,
    heading: bool = False,
    align=None,
) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    elif heading:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    _set_run_font(run, size, bold=bold)


def _set_title_paragraph(paragraph, topic: str) -> None:
    """纲目主题：三号 16pt、加粗、居中（复用 _set_run_font）。"""
    for child in list(paragraph._element):
        # 保留 pPr，只清 run
        if child.tag.endswith("}r"):
            paragraph._element.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(topic or "")
    _set_run_font(run, 16, bold=True)


def _cost_line(run: dict) -> str:
    sel = run.get("cost_select_usd")
    ref = run.get("cost_ref_usd")
    plain = run.get("cost_plain_usd")
    total = run.get("cost_total_usd")
    from_cache = bool(run.get("plain_from_cache"))
    if plain is None and ref is None and sel is None and total is None:
        return "费用：—"
    plain_part = "存档" if from_cache else format_usd(plain)
    return (
        f"费用：选词 {format_usd(sel)}｜加词条定义版 {format_usd(ref)}"
        f"｜普通版 {plain_part}｜合计 {format_usd(total)}"
    )


def extract_burden_export(text: str) -> tuple[str, str | None]:
    """取【负担说明】之后正文，到下一个以「【」开头的标题行之前；找不到则全文 + warning。"""
    raw = text or ""
    index = raw.rfind(BURDEN_MARKER)
    if index < 0:
        return raw, "未找到【负担说明】标题，已导出全文"
    rest = raw[index + len(BURDEN_MARKER) :]
    lead = rest.lstrip(" \t\r\n")
    if lead.startswith(BURDEN_TAG):
        rest = lead[len(BURDEN_TAG) :]
    else:
        m = _HINT_AFTER_BURDEN.match(lead)
        if m:
            rest = lead[m.end() :]
        else:
            rest = lead
    rest = rest.lstrip("\r\n")
    # 截到下一个【标题行】之前（直写版末尾的【用到的参考材料】不导出）
    out_lines: list[str] = []
    for line in rest.splitlines(keepends=True):
        stripped = line.lstrip("\r\n")
        if out_lines and stripped.startswith("【"):
            break
        out_lines.append(line)
    body = "".join(out_lines).rstrip("\r\n")
    return body, None


def kind_docx_filename(kind: ExportKind, topic: str) -> str:
    label = _KIND_LABEL[kind]
    prefix = f"【{label}】"
    # 截断只截纲目主题，不截前缀
    max_topic = max(1, 100 - len(prefix) - len(".docx"))
    safe = safe_filename_part(topic, max_topic)
    return f"{prefix}{safe}.docx"


def kind_title_text(kind: ExportKind, topic: str) -> str:
    """docx 首段文字：选词结果仍为纲目主题；另两种带【…-负担说明】前缀。"""
    t = (topic or "").strip()
    if kind == "select":
        return t
    return f"【{_KIND_LABEL[kind]}】{t}"


def build_kind_docx(kind: ExportKind, topic: str, text: str) -> tuple[bytes, str | None]:
    """三种独立导出。返回 (docx_bytes, warning|None)。"""
    warning = None
    if kind == "select":
        body = text or ""
    else:
        body, warning = extract_burden_export(text or "")

    title = kind_title_text(kind, topic)
    doc = Document()
    if doc.paragraphs:
        _set_title_paragraph(doc.paragraphs[0], title)
    else:
        p = doc.add_paragraph()
        _set_title_paragraph(p, title)

    lines = body.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if body == "":
        lines = []
    for line in lines:
        _add(doc, line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), warning


def export_warning_header(warning: str | None) -> dict[str, str]:
    if not warning:
        return {}
    return {"X-SC-Export-Warning": quote(warning)}


def build_sc_docx(run: dict) -> bytes:
    doc = Document()
    topic = run.get("topic") or ""
    _add(doc, _cost_line(run), bold=True, size=11)
    _add(doc, f"篇题：{topic}", bold=True, size=14, heading=True)

    _add(doc, "最终选词名单", bold=True, size=12, heading=True)
    terms = run.get("final_terms") or []
    if isinstance(terms, str):
        terms = json.loads(terms)
    for item in terms:
        if isinstance(item, str):
            _add(doc, item)
        else:
            name = item.get("name", "")
            tier = item.get("tier") or item.get("tier_actual") or ""
            reason = item.get("reason") or ""
            line = f"〔{tier}〕{name}"
            if reason:
                line += f"｜依据：{reason}"
            _add(doc, line)

    skipped = run.get("skipped") or run.get("skipped_json") or []
    merged = run.get("merged") or run.get("merged_json") or []
    if isinstance(skipped, str):
        skipped = json.loads(skipped or "[]")
    if isinstance(merged, str):
        merged = json.loads(merged or "[]")
    _add(doc, "被跳过（无 3a）", bold=True, size=12, heading=True)
    _add(doc, "、".join(skipped) if skipped else "（无）")
    _add(doc, "被合并（定义相同）", bold=True, size=12, heading=True)
    _add(doc, "、".join(merged) if merged else "（无）")

    _add(doc, "普通版", bold=True, size=12, heading=True)
    plain = run.get("plain_result") or ""
    for line in (plain or "（无）").splitlines() or ["（无）"]:
        _add(doc, line)

    _add(doc, "加词条定义版", bold=True, size=12, heading=True)
    ref = run.get("ref_result") or ""
    for line in (ref or "（无）").splitlines() or ["（无）"]:
        _add(doc, line)

    _append_citation_page(doc, run)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _fmt_items(items) -> str:
    if not items:
        return "—"
    return "、".join(str(x) for x in items)


def _append_citation_page(doc: Document, run: dict) -> None:
    stats = run.get("citation_stats")
    if stats is None:
        stats = compute_citation_stats(
            ref_output=run.get("ref_result"),
            final_terms=run.get("final_terms") or [],
            skipped=run.get("skipped") or run.get("skipped_json") or [],
            merged=run.get("merged") or run.get("merged_json") or [],
        )
    # 分页
    p = doc.add_paragraph()
    run_br = p.add_run()
    run_br.add_break(WD_BREAK.PAGE)

    _add(doc, "引用统计", bold=True, size=14, heading=True)
    if stats.get("note"):
        _add(doc, stats["note"])
        return

    if stats.get("format_note"):
        _add(doc, stats["format_note"])

    by_tier = stats.get("by_tier") or {}
    fmt = stats.get("format") or "outline"
    if fmt == "direct":
        listed = sum(
            (by_tier.get(t, {}).get("listed", 0) + by_tier.get(t, {}).get("both", 0))
            for t in TIER_ORDER
        )
        in_text = sum(by_tier.get(t, {}).get("in_text", 0) for t in TIER_ORDER)
        unc = sum(by_tier.get(t, {}).get("unused", 0) for t in TIER_ORDER)
        _add(
            doc,
            f"基数 {stats.get('sent_total', 0)}；已列出 {listed}；仅文中出现 {in_text}；未使用 {unc}",
            bold=True,
            size=11,
        )
        col_header = "词条名｜状态"
    else:
        in_o = sum(by_tier.get(t, {}).get("in_outline", 0) for t in TIER_ORDER)
        defs_o = sum(by_tier.get(t, {}).get("defs_only", 0) for t in TIER_ORDER)
        unc = sum(by_tier.get(t, {}).get("uncited", 0) for t in TIER_ORDER)
        _add(
            doc,
            f"基数 {stats.get('sent_total', 0)}；入脉络 {in_o}；仅界定 {defs_o}；未引用 {unc}",
            bold=True,
            size=11,
        )
        col_header = "词条名｜状态｜职事界定序号｜真理脉络环序号"

    terms = stats.get("terms") or []
    for tier in TIER_ORDER:
        _add(doc, f"〔{tier}〕", bold=True, size=12, heading=True)
        _add(doc, col_header, bold=True, size=11)
        rows = [t for t in terms if t.get("tier") == tier]
        if not rows:
            _add(doc, "（无）")
            continue
        for t in rows:
            label = t.get("status_label") or STATUS_LABEL.get(t.get("status") or "", t.get("status") or "")
            if fmt == "direct":
                _add(doc, f"{t.get('name')}｜{label}")
            else:
                _add(
                    doc,
                    f"{t.get('name')}｜{label}｜{_fmt_items(t.get('def_items'))}｜{_fmt_items(t.get('outline_items'))}",
                )

    _add(doc, "引用了未选词条", bold=True, size=12, heading=True)
    unknown = stats.get("unknown_cited") or []
    if not unknown:
        _add(doc, "无")
    else:
        for u in unknown:
            if fmt == "direct":
                _add(doc, str(u.get("name") or ""))
            else:
                _add(
                    doc,
                    f"{u.get('name')}｜界定 {_fmt_items(u.get('def_items'))}｜脉络 {_fmt_items(u.get('outline_items'))}",
                )


def sc_docx_filename(topic: str, run_id: int) -> str:
    return f"SC_{safe_filename_part(topic, 30)}_{run_id}.docx"
