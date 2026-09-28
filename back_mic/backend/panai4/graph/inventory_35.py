# -*- coding: utf-8 -*-
"""盘点 3.5 Neo4j（只读）、核对 panai4 词表资料，并写出对应表草稿。

连接方式与 kg_rag/neo4j_client.py、kg_rag/scripts/import_concepts.py 相同：
从 back_mic/backend/.env 读取 NEO4J_URI / NEO4J_USER / NEO4J_PASSWORD。
密码不写入日志或报告。所有 Cypher 都经 execute_read 提交，且以只读关键字开头。

用法（在仓库任意目录）:
    python back_mic/backend/panai4/graph/inventory_35.py
"""
from __future__ import annotations

import csv
import json
import os
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from neo4j import READ_ACCESS, GraphDatabase

_BACKEND = Path(__file__).resolve().parents[2]
_REPO = _BACKEND.parents[1]
_DATA = _REPO / "docs" / "panai4_graph"
_REPORTS = _DATA / "reports"

_TXT = _DATA / "2372个词条（已确定一一对应为concept节点）.txt"
_TERMS_JSON = _DATA / "主恢复真理词典_词条主干枝子映射 (1).json"
_TAGS_JSON = _DATA / "主恢复真理词典_标签映射_已校对 (1).json"
_TRUNKS_TXT = _DATA / "12大类与91小类的对应关系.txt"
_DOCX = _DATA / "Neo4j节点与真理词典比对报告.docx"

_PAREN_RE = re.compile(r"^([^（）()]+)（([^（）()]+)）$")
_ANY_PAREN_RE = re.compile(r"[（）()]")
_WRITE_RE = re.compile(
    r"^\s*(CREATE|MERGE|SET|DELETE|DETACH|DROP|REMOVE|INSERT|LOAD\s+CSV)\b",
    re.IGNORECASE,
)


def _read(tx, query: str, params: dict | None = None):
    if _WRITE_RE.search(query):
        raise RuntimeError(f"拒绝执行非只读语句: {query[:80]}")
    result = tx.run(query, params or {})
    return [record.data() for record in result]


def _run(session, query: str, params: dict | None = None):
    return session.execute_read(_read, query, params)


def _load_env() -> dict[str, str]:
    env_path = _BACKEND / ".env"
    if env_path.is_file():
        load_dotenv(env_path)
    uri = os.environ.get("NEO4J_URI", "bolt://localhost:7687")
    user = os.environ.get("NEO4J_USER", "neo4j")
    password = os.environ.get("NEO4J_PASSWORD", "")
    return {"uri": uri, "user": user, "password": password}


def _redact(text: str, password: str) -> str:
    if password and password in text:
        return text.replace(password, "***")
    return text


def _nonempty(value) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip() != ""
    if isinstance(value, (list, tuple, dict)):
        return len(value) > 0
    return True


def _label_key(labels) -> str:
    return "+".join(sorted(labels or [])) or "(无标签)"


# ---------------------------------------------------------------------------
# 资料
# ---------------------------------------------------------------------------

def load_terms_txt() -> tuple[list[str], list[str]]:
    """返回 (去空行并 strip 后的词条, 原始非空行 strip 前有首尾空白的原文)。"""
    raw = _TXT.read_text(encoding="utf-8-sig")
    names: list[str] = []
    dirty: list[str] = []
    for line in raw.splitlines():
        if line.strip() == "":
            continue
        if line != line.strip():
            dirty.append(line)
        names.append(line.strip())
    return names, dirty


def parse_trunks(text: str) -> tuple[list[tuple[str, list[str]]], list[str]]:
    """解析 12 大类文件。返回 [(主干, [枝子...])], 解析提示。"""
    notes: list[str] = []
    trunks: list[tuple[str, list[str]]] = []
    current: str | None = None
    expected_n: int | None = None
    branches: list[str] = []
    header_re = re.compile(r"^主干分类第(\d+)大类：(.+)$")
    tag_re = re.compile(r"^横切标签（(\d+)小类）：$")

    def flush():
        nonlocal current, expected_n, branches
        if current is None:
            return
        if expected_n is not None and len(branches) != expected_n:
            notes.append(
                f"主干「{current}」标题写 {expected_n} 个枝子，实际读到 {len(branches)} 个"
            )
        trunks.append((current, branches))
        current = None
        expected_n = None
        branches = []

    for line in text.splitlines():
        stripped = line.strip()
        m = header_re.match(stripped)
        if m:
            flush()
            current = m.group(2).strip()
            continue
        m = tag_re.match(stripped)
        if m and current is not None:
            expected_n = int(m.group(1))
            continue
        if current is not None and stripped and not stripped.startswith("主干分类共"):
            if stripped.startswith("以下是"):
                continue
            branches.append(stripped)
    flush()
    return trunks, notes


def load_docx() -> tuple[list[tuple[str, str]], list[str], dict]:
    """同义表 (旧名, 新名) 与找不到表旧名。第三项是文档自述数量。"""
    doc = Document(str(_DOCX))
    stated = {}
    if len(doc.tables) < 3:
        raise RuntimeError(f"比对报告表格数量为 {len(doc.tables)}，期望至少 3")
    summary = doc.tables[0]
    for row in summary.rows[1:]:
        k = row.cells[0].text.strip()
        v = row.cells[1].text.strip()
        stated[k] = v
    synonyms: list[tuple[str, str]] = []
    for row in doc.tables[1].rows[1:]:
        old = row.cells[0].text.strip()
        new = row.cells[1].text.strip()
        if old or new:
            synonyms.append((old, new))
    missing: list[str] = []
    for row in doc.tables[2].rows[1:]:
        left = row.cells[1].text.strip()
        right = row.cells[3].text.strip()
        if left:
            missing.append(left)
        if right:
            missing.append(right)
    return synonyms, missing, stated


def build_2380(terms: list[str]) -> tuple[list[str], list[tuple[str, str, str]], list[str]]:
    """返回 (2380 名单, [(原词, 括号外, 括号内)], 无法按单层括号拆开的行)。"""
    splits: list[tuple[str, str, str]] = []
    bad: list[str] = []
    out: list[str] = []
    for name in terms:
        if _ANY_PAREN_RE.search(name):
            m = _PAREN_RE.match(name)
            if not m:
                bad.append(name)
                out.append(name)
                continue
            outer, inner = m.group(1).strip(), m.group(2).strip()
            splits.append((name, outer, inner))
            out.append(outer)
            out.append(inner)
        else:
            out.append(name)
    return out, splits, bad


# ---------------------------------------------------------------------------
# Neo4j 只读盘点
# ---------------------------------------------------------------------------

def inventory_graph(session) -> dict:
    labels = _run(
        session,
        "MATCH (n) RETURN labels(n) AS labels, count(*) AS cnt ORDER BY cnt DESC",
    )
    rels = _run(
        session,
        """
        MATCH (a)-[r]->(b)
        RETURN type(r) AS rel_type, labels(a) AS start_labels, labels(b) AS end_labels,
               count(*) AS cnt
        ORDER BY rel_type, cnt DESC
        """,
    )
    rel_keys = _run(
        session,
        """
        MATCH ()-[r]->()
        WITH type(r) AS rel_type, keys(r) AS ks
        UNWIND CASE WHEN size(ks)=0 THEN [null] ELSE ks END AS k
        RETURN rel_type, k AS prop, count(*) AS cnt
        ORDER BY rel_type, prop
        """,
    )
    concepts = _run(session, "MATCH (c:Concept) RETURN properties(c) AS props")
    scriptures = _run(session, "MATCH (s:Scripture) RETURN s.id AS id, keys(s) AS ks")
    scripture_key_stats = _run(
        session,
        """
        MATCH (s:Scripture)
        UNWIND keys(s) AS k
        WITH k, s[k] AS v
        RETURN k AS prop,
               count(*) AS present,
               sum(CASE
                   WHEN v IS NULL THEN 0
                   WHEN v = '' THEN 0
                   ELSE 1 END) AS nonempty_scalar
        ORDER BY prop
        """,
    )
    constraints = _run(session, "SHOW CONSTRAINTS")
    indexes = _run(session, "SHOW INDEXES")
    edges = _run(
        session,
        """
        MATCH (a)-[r]->(b)
        RETURN type(r) AS rel_type,
               labels(a) AS start_labels,
               labels(b) AS end_labels,
               CASE WHEN a:Concept THEN a.name ELSE null END AS start_name,
               CASE WHEN a:Scripture THEN a.id ELSE null END AS start_id,
               CASE WHEN b:Concept THEN b.name ELSE null END AS end_name,
               CASE WHEN b:Scripture THEN b.id ELSE null END AS end_id
        """,
    )
    return {
        "labels": labels,
        "rels": rels,
        "rel_keys": rel_keys,
        "concepts": concepts,
        "scriptures": scriptures,
        "scripture_key_stats": scripture_key_stats,
        "constraints": constraints,
        "indexes": indexes,
        "edges": edges,
    }


def _prop_stats(prop_dicts: list[dict]) -> list[tuple[str, int, int]]:
    present: Counter[str] = Counter()
    nonempty: Counter[str] = Counter()
    for props in prop_dicts:
        for key, value in props.items():
            present[key] += 1
            if _nonempty(value):
                nonempty[key] += 1
    rows = []
    for key in sorted(present):
        rows.append((key, present[key], nonempty[key]))
    return rows


def _name_anomalies(names: list[str]) -> dict:
    counter = Counter(names)
    duplicates = sorted((n, c) for n, c in counter.items() if c > 1)
    nulls = sum(1 for n in names if n is None)
    blanks = [n for n in names if n is not None and str(n).strip() == ""]
    padded = []
    for n in names:
        if not isinstance(n, str):
            continue
        if n != n.strip() or n != n.strip("\u3000"):
            padded.append(n)
    half = [n for n in names if isinstance(n, str) and (("(" in n) or (")" in n))]
    full = [n for n in names if isinstance(n, str) and (("（" in n) or ("）" in n))]
    mixed = [
        n
        for n in names
        if isinstance(n, str) and (("(" in n) or (")" in n)) and (("（" in n) or ("）" in n))
    ]
    inner_space = [
        n
        for n in names
        if isinstance(n, str) and (("  " in n) or ("\u3000" in n))
    ]
    return {
        "duplicates": duplicates,
        "nulls": nulls,
        "blanks": blanks,
        "padded": sorted(set(padded)),
        "half": sorted(set(half)),
        "full": sorted(set(full)),
        "mixed": sorted(set(mixed)),
        "inner_space": sorted(set(inner_space)),
        "distinct": len(counter),
        "total": len(names),
    }


# ---------------------------------------------------------------------------
# 对应表
# ---------------------------------------------------------------------------

def build_mapping(
    concept_names: list[str],
    list_2380: set[str],
    splits: list[tuple[str, str, str]],
    synonyms: list[tuple[str, str]],
    missing: list[str],
) -> tuple[list[dict], dict]:
    paren_full = {src for src, _, _ in splits}
    split_parts = {part for _, outer, inner in splits for part in (outer, inner)}
    syn_map = {old: new for old, new in synonyms}
    missing_set = set(missing)

    rule_b_olds = {old for old, new in synonyms if new in paren_full}
    rule_c_olds = {name for name in missing if name in split_parts}

    rows: list[dict] = []
    seen: set[str] = set()
    for old in concept_names:
        if old in seen:
            continue
        seen.add(old)
        if old in rule_c_olds or (old in missing_set and old in split_parts):
            rows.append(
                {
                    "旧名": old,
                    "新名": old,
                    "处理方式": "一致",
                    "备注": "比对报告「完全找不到」表中的名称，与某个括号词条拆出的名字相同，按规则视为按名直接对应",
                }
            )
        elif old in missing_set:
            note = "比对报告「完全找不到」，不迁移"
            if old in list_2380:
                note += "；同时这个旧名也在 2380 名单中，与「删除」规则冲突，需人工决定"
            rows.append({"旧名": old, "新名": "", "处理方式": "删除", "备注": note})
        elif old in rule_b_olds or (old in syn_map and syn_map[old] in paren_full):
            rows.append(
                {
                    "旧名": old,
                    "新名": old,
                    "处理方式": "一致",
                    "备注": f"比对报告「同义」表的新名是带括号词条「{syn_map.get(old, '')}」，拆开后与旧名同名，按规则视为按名直接对应",
                }
            )
        elif old in syn_map:
            new = syn_map[old]
            note = "比对报告「同义但说法不同」"
            if old in list_2380:
                note += "；旧名本身已在 2380 名单中，与「改名」并存，需人工决定"
            if new not in list_2380:
                note += "；新名不在 2380 名单中"
            rows.append({"旧名": old, "新名": new, "处理方式": "改名", "备注": note})
        elif old in list_2380:
            rows.append({"旧名": old, "新名": old, "处理方式": "一致", "备注": ""})
        else:
            rows.append(
                {
                    "旧名": old,
                    "新名": "",
                    "处理方式": "未归类",
                    "备注": "3.5 中存在，但不在 2380 名单，也不在比对报告的两张表里",
                }
            )
    rows.sort(key=lambda r: (r["旧名"] or ""))

    # 报告里有、图里没有
    graph_names = set(concept_names)
    report_only = []
    for old, new in synonyms:
        if old not in graph_names:
            report_only.append(("同义表", old, new))
    for name in missing:
        if name not in graph_names:
            report_only.append(("找不到表", name, ""))

    new_name_owners: dict[str, list[str]] = defaultdict(list)
    for row in rows:
        if row["处理方式"] in ("一致", "改名") and row["新名"]:
            new_name_owners[row["新名"]].append(row["旧名"])
    collisions = {k: v for k, v in new_name_owners.items() if len(v) > 1}

    counts = Counter(r["处理方式"] for r in rows)
    renamed_not_in_2380 = [
        r for r in rows if r["处理方式"] == "改名" and r["新名"] not in list_2380
    ]
    deleted_but_in_2380 = [
        r for r in rows if r["处理方式"] == "删除" and r["旧名"] in list_2380
    ]
    meta = {
        "counts": counts,
        "rule_b": sorted(rule_b_olds),
        "rule_c": sorted(rule_c_olds),
        "report_only": report_only,
        "collisions": collisions,
        "renamed_not_in_2380": renamed_not_in_2380,
        "deleted_but_in_2380": deleted_but_in_2380,
        "syn_duplicate_olds": [n for n, c in Counter(o for o, _ in synonyms).items() if c > 1],
        "missing_duplicate": [n for n, c in Counter(missing).items() if c > 1],
        "paren_full": sorted(paren_full),
        "split_parts": split_parts,
    }
    return rows, meta


def migration_estimate(edges: list[dict], rows: list[dict]) -> list[dict]:
    by_old = {r["旧名"]: r for r in rows}

    def map_end(labels, name, sid):
        labs = set(labels or [])
        if "Concept" in labs:
            row = by_old.get(name)
            if row is None:
                return ("unmapped", None)
            if row["处理方式"] == "删除":
                return ("deleted", None)
            if row["处理方式"] == "未归类" or not row["新名"]:
                return ("unmapped", None)
            return ("ok", ("Concept", row["新名"]))
        if "Scripture" in labs:
            if sid is None or str(sid).strip() == "":
                return ("unmapped", None)
            return ("ok", ("Scripture", str(sid)))
        return ("unmapped", None)

    grouped: dict[str, list[dict]] = defaultdict(list)
    for edge in edges:
        grouped[edge["rel_type"]].append(edge)

    stats = []
    for rel_type in sorted(grouped):
        items = grouped[rel_type]
        both = 0
        discarded = 0
        unmapped = 0
        pair_counter: Counter[tuple] = Counter()
        for edge in items:
            sk, skey = map_end(edge["start_labels"], edge["start_name"], edge["start_id"])
            ek, ekey = map_end(edge["end_labels"], edge["end_name"], edge["end_id"])
            if sk == "deleted" or ek == "deleted":
                discarded += 1
                continue
            if sk != "ok" or ek != "ok":
                unmapped += 1
                continue
            both += 1
            pair_counter[(skey, ekey)] += 1
        dup_groups = [(k, c) for k, c in pair_counter.items() if c > 1]
        extra = sum(c - 1 for _, c in dup_groups)
        examples = []
        for (skey, ekey), c in sorted(dup_groups, key=lambda x: -x[1])[:15]:
            examples.append(f"{skey[0]}:{skey[1]} → {ekey[0]}:{ekey[1]} ×{c}")
        stats.append(
            {
                "rel_type": rel_type,
                "total": len(items),
                "both": both,
                "discarded": discarded,
                "unmapped": unmapped,
                "dup_groups": len(dup_groups),
                "extra": extra,
                "examples": examples,
            }
        )
    return stats


# ---------------------------------------------------------------------------
# 资料核对
# ---------------------------------------------------------------------------

def cross_check(terms: list[str], term_rows: list[dict], tag_obj: dict, trunk_pairs: list[tuple[str, list[str]]]) -> dict:
    term_set = set(terms)
    json_names = [str(r.get("name", "")).strip() for r in term_rows]
    json_set = set(json_names)
    trunk_names = [t for t, _ in trunk_pairs]
    trunk_set = set(trunk_names)
    branch_to_trunk: dict[str, str] = {}
    branch_dupes: list[tuple[str, str, str]] = []
    for trunk, branches in trunk_pairs:
        for b in branches:
            if b in branch_to_trunk:
                branch_dupes.append((b, branch_to_trunk[b], trunk))
            else:
                branch_to_trunk[b] = trunk
    branch_set = set(branch_to_trunk)

    only_txt = sorted(term_set - json_set)
    only_json = sorted(json_set - term_set)
    json_dupes = sorted(n for n, c in Counter(json_names).items() if c > 1)

    trunk_bad = []
    direct_not_branch = []
    direct_trunk_mismatch = []
    indirect_not_branch = []
    for r in term_rows:
        name = str(r.get("name", "")).strip()
        trunk = str(r.get("trunk", "")).strip()
        direct = str(r.get("branch_directly", "")).strip()
        if trunk not in trunk_set:
            trunk_bad.append((name, trunk))
        if direct not in branch_set:
            direct_not_branch.append((name, direct))
        elif branch_to_trunk[direct] != trunk:
            direct_trunk_mismatch.append((name, trunk, direct, branch_to_trunk[direct]))
        for ind in r.get("branch_indirectly") or []:
            ind_s = str(ind).strip()
            if ind_s not in branch_set:
                indirect_not_branch.append((name, ind_s))

    # 主标签 / 副标签
    direct_members: dict[str, set[str]] = defaultdict(set)
    indirect_members: dict[str, set[str]] = defaultdict(set)
    for r in term_rows:
        name = str(r.get("name", "")).strip()
        direct = str(r.get("branch_directly", "")).strip()
        direct_members[direct].add(name)
        for ind in r.get("branch_indirectly") or []:
            indirect_members[str(ind).strip()].add(name)

    tag_keys = list(tag_obj.keys())
    tag_trunk_mismatch = []
    tag_only = sorted(set(tag_keys) - branch_set)
    branch_only = sorted(branch_set - set(tag_keys))
    for key, payload in tag_obj.items():
        stated = str((payload or {}).get("主干分类", "")).strip()
        if key in branch_to_trunk and stated != branch_to_trunk[key]:
            tag_trunk_mismatch.append((key, stated, branch_to_trunk[key]))
        elif key not in branch_to_trunk:
            pass

    primary_mismatch = []
    secondary_mismatch = []
    primary_internal_dupes = []
    secondary_internal_dupes = []
    for key in sorted(set(tag_keys) | branch_set):
        payload = tag_obj.get(key) or {}
        primary_list = [str(x).strip() for x in (payload.get("主标签") or [])]
        secondary_list = [str(x).strip() for x in (payload.get("副标签") or [])]
        if len(primary_list) != len(set(primary_list)):
            primary_internal_dupes.append((key, [n for n, c in Counter(primary_list).items() if c > 1]))
        if len(secondary_list) != len(set(secondary_list)):
            secondary_internal_dupes.append((key, [n for n, c in Counter(secondary_list).items() if c > 1]))
        p_set = set(primary_list)
        s_set = set(secondary_list)
        d_set = direct_members.get(key, set())
        i_set = indirect_members.get(key, set())
        if p_set != d_set:
            primary_mismatch.append(
                {
                    "枝子": key,
                    "只在主标签": sorted(p_set - d_set),
                    "只在 branch_directly": sorted(d_set - p_set),
                }
            )
        if s_set != i_set:
            secondary_mismatch.append(
                {
                    "枝子": key,
                    "只在副标签": sorted(s_set - i_set),
                    "只在 branch_indirectly 反推": sorted(i_set - s_set),
                }
            )

    return {
        "trunk_names": trunk_names,
        "trunk_set": trunk_set,
        "branch_to_trunk": branch_to_trunk,
        "branch_dupes": branch_dupes,
        "only_txt": only_txt,
        "only_json": only_json,
        "json_dupes": json_dupes,
        "json_count": len(json_names),
        "trunk_bad": trunk_bad,
        "direct_not_branch": direct_not_branch,
        "direct_trunk_mismatch": direct_trunk_mismatch,
        "indirect_not_branch": indirect_not_branch,
        "tag_key_count": len(tag_keys),
        "tag_only": tag_only,
        "branch_only": branch_only,
        "tag_trunk_mismatch": tag_trunk_mismatch,
        "primary_mismatch": primary_mismatch,
        "secondary_mismatch": secondary_mismatch,
        "primary_internal_dupes": primary_internal_dupes,
        "secondary_internal_dupes": secondary_internal_dupes,
        "tag_key_dupes": [n for n, c in Counter(tag_keys).items() if c > 1],
    }


# ---------------------------------------------------------------------------
# 写报告
# ---------------------------------------------------------------------------

def _yn(ok: bool) -> str:
    return "通过" if ok else "不通过"


def _bullet_list(items: list, limit: int | None = None) -> list[str]:
    shown = items if limit is None else items[:limit]
    lines = [f"- {x}" for x in shown]
    if limit is not None and len(items) > limit:
        lines.append(f"- …其余 {len(items) - limit} 条未展开")
    if not lines:
        lines.append("- （无）")
    return lines


def write_source_check(
    path: Path,
    terms: list[str],
    dirty_lines: list[str],
    term_dupes: list[str],
    splits: list[tuple[str, str, str]],
    bad_parens: list[str],
    list_2380: list[str],
    collisions_2380: list[tuple[str, list[str]]],
    trunk_pairs: list[tuple[str, list[str]]],
    trunk_notes: list[str],
    cross: dict,
    synonyms: list[tuple[str, str]],
    missing: list[str],
    stated: dict,
    map_rows: list[dict],
    map_meta: dict,
    concept_node_count: int,
) -> None:
    lines: list[str] = []
    a = lines.append
    a("# 资料核对")
    a("")
    a("本报告由 `back_mic/backend/panai4/graph/inventory_35.py` 生成。只读，未改词表，也未写 Neo4j。")
    a("")
    a("## 1. 2372 词表 txt")
    a("")
    term_counter = Counter(terms)
    dupes = sorted(n for n, c in term_counter.items() if c > 1)
    ok_count = len(terms) == 2372 and not dupes
    a(f"- 去掉空行并去掉首尾空格后的行数：**{len(terms)}**（期望 2372）→ {_yn(len(terms) == 2372)}")
    a(f"- 重复词条：**{len(dupes)}** → {_yn(not dupes)}")
    if dupes:
        lines.extend(_bullet_list([f"{n} ×{term_counter[n]}" for n in dupes]))
    a(f"- 原文行有首尾空白（strip 之后才入库的行）：**{len(dirty_lines)}**")
    if dirty_lines:
        lines.extend(_bullet_list([repr(x) for x in dirty_lines]))
    a(f"- 带括号、且能拆成「括号外 + 括号内」的词条：**{len(splits)}**（期望 8）→ {_yn(len(splits) == 8)}")
    for src, outer, inner in splits:
        a(f"  - `{src}` → `{outer}` + `{inner}`")
    if bad_parens:
        a(f"- 含括号但不是「单层全角一对」的行：**{len(bad_parens)}**")
        lines.extend(_bullet_list(bad_parens))
    a("")
    a("## 2. 12 大类与 91 小类")
    a("")
    branch_list = [b for _, bs in trunk_pairs for b in bs]
    branch_counter = Counter(branch_list)
    branch_dupes = sorted(n for n, c in branch_counter.items() if c > 1)
    a(f"- 主干数：**{len(trunk_pairs)}**（期望 12）→ {_yn(len(trunk_pairs) == 12)}")
    a(f"- 枝子数：**{len(branch_list)}**（期望 91）→ {_yn(len(branch_list) == 91)}")
    a(f"- 枝子重复：**{len(branch_dupes)}** → {_yn(not branch_dupes)}")
    if branch_dupes:
        lines.extend(_bullet_list([f"{n} ×{branch_counter[n]}" for n in branch_dupes]))
    if cross["branch_dupes"]:
        a("- 同一枝子出现在不同主干下：")
        for b, t1, t2 in cross["branch_dupes"]:
            a(f"  - `{b}`：`{t1}` 与 `{t2}`")
    if trunk_notes:
        a("- 标题里的小类个数与实际行数不一致：")
        lines.extend(_bullet_list(trunk_notes))
    a("")
    a("主干与枝子：")
    a("")
    for trunk, branches in trunk_pairs:
        a(f"- **{trunk}**（{len(branches)}）：{'、'.join(branches)}")
    a("")
    a("## 3. 两份 JSON 与 12 大类、txt 的交叉核对")
    a("")
    a("### 3.1 词条集合")
    a("")
    a(f"- txt 词条数：**{len(terms)}**（去重后 {len(set(terms))}）")
    a(f"- 词条映射 JSON 条目数：**{cross['json_count']}**，其中重复 name：**{len(cross['json_dupes'])}**")
    if cross["json_dupes"]:
        lines.extend(_bullet_list(cross["json_dupes"]))
    a(f"- 与 txt 的集合是否完全一致：{_yn(not cross['only_txt'] and not cross['only_json'])}")
    a(f"- 只在 txt 中：**{len(cross['only_txt'])}**")
    lines.extend(_bullet_list(cross["only_txt"]))
    a(f"- 只在词条映射 JSON 中：**{len(cross['only_json'])}**")
    lines.extend(_bullet_list(cross["only_json"]))
    a("")
    a("### 3.2 每个词条的 trunk")
    a("")
    a(f"- trunk 不属于 12 个主干的词条：**{len(cross['trunk_bad'])}** → {_yn(not cross['trunk_bad'])}")
    lines.extend(_bullet_list([f"`{n}` → trunk=`{t}`" for n, t in cross["trunk_bad"]]))
    a("")
    a("### 3.3 branch_directly")
    a("")
    a(f"- 不属于 91 个枝子：**{len(cross['direct_not_branch'])}** → {_yn(not cross['direct_not_branch'])}")
    lines.extend(_bullet_list([f"`{n}` → `{b}`" for n, b in cross["direct_not_branch"]]))
    a(f"- 枝子在 12 大类中的主干 ≠ 该词条的 trunk：**{len(cross['direct_trunk_mismatch'])}** → {_yn(not cross['direct_trunk_mismatch'])}")
    lines.extend(
        _bullet_list(
            [
                f"`{n}`：词条 trunk=`{t}`，branch_directly=`{b}`，该枝子归属主干=`{tb}`"
                for n, t, b, tb in cross["direct_trunk_mismatch"]
            ]
        )
    )
    a("")
    a("### 3.4 branch_indirectly")
    a("")
    a(f"- 值不属于 91 个枝子的次数：**{len(cross['indirect_not_branch'])}** → {_yn(not cross['indirect_not_branch'])}")
    lines.extend(_bullet_list([f"`{n}` → `{b}`" for n, b in cross["indirect_not_branch"]]))
    a("")
    a("### 3.5 标签映射 JSON 与 12 大类")
    a("")
    a(f"- 标签映射键数：**{cross['tag_key_count']}**（期望 91）→ {_yn(cross['tag_key_count'] == 91 and not cross['tag_only'] and not cross['branch_only'])}")
    a(f"- 键有、12 大类没有：**{len(cross['tag_only'])}**")
    lines.extend(_bullet_list(cross["tag_only"]))
    a(f"- 12 大类有、键没有：**{len(cross['branch_only'])}**")
    lines.extend(_bullet_list(cross["branch_only"]))
    a(f"- 「主干分类」与 12 大类不一致：**{len(cross['tag_trunk_mismatch'])}** → {_yn(not cross['tag_trunk_mismatch'])}")
    lines.extend(
        _bullet_list(
            [f"`{k}`：标签映射写 `{got}`，12 大类为 `{exp}`" for k, got, exp in cross["tag_trunk_mismatch"]]
        )
    )
    a("")
    a("### 3.6 主标签 vs branch_directly")
    a("")
    a("比较的是集合，不计顺序。")
    a(f"- 不一致的枝子：**{len(cross['primary_mismatch'])}** → {_yn(not cross['primary_mismatch'])}")
    a(f"- 主标签列表内部有重复的枝子：**{len(cross['primary_internal_dupes'])}**")
    for key, dups in cross["primary_internal_dupes"]:
        a(f"  - `{key}`：{ '、'.join(dups) }")
    for item in cross["primary_mismatch"]:
        a(f"- 枝子 `{item['枝子']}`")
        a(f"  - 只在主标签（{len(item['只在主标签'])}）：{'、'.join(item['只在主标签']) or '（无）'}")
        a(f"  - 只在 branch_directly（{len(item['只在 branch_directly'])}）：{'、'.join(item['只在 branch_directly']) or '（无）'}")
    a("")
    a("### 3.7 副标签 vs branch_indirectly 反推")
    a("")
    a("反推：词条映射里 `branch_indirectly` 包含该枝子的词条集合。未做任何修正。")
    a(f"- 不一致的枝子：**{len(cross['secondary_mismatch'])}** → {_yn(not cross['secondary_mismatch'])}")
    a(f"- 副标签列表内部有重复的枝子：**{len(cross['secondary_internal_dupes'])}**")
    for key, dups in cross["secondary_internal_dupes"]:
        a(f"  - `{key}`：{'、'.join(dups)}")
    for item in cross["secondary_mismatch"]:
        a(f"- 枝子 `{item['枝子']}`")
        a(f"  - 只在副标签（{len(item['只在副标签'])}）：{'、'.join(item['只在副标签']) or '（无）'}")
        a(
            f"  - 只在 branch_indirectly 反推（{len(item['只在 branch_indirectly 反推'])}）："
            f"{'、'.join(item['只在 branch_indirectly 反推']) or '（无）'}"
        )
    a("")
    a("## 4. 2380 名单")
    a("")
    a("规则 A：8 个带括号词条各拆成两个名字，括号本身不保留。2380 = 2372 − 8 + 16。")
    unique_2380 = len(set(list_2380))
    a(f"- 拆分后名单长度：**{len(list_2380)}**（期望 2380）→ {_yn(len(list_2380) == 2380)}")
    a(f"- 去重后：**{unique_2380}** → {_yn(unique_2380 == len(list_2380) == 2380)}")
    a(f"- 拆出的名字与原词表已有词条重名，或两个拆出名字彼此重名：**{len(collisions_2380)}**")
    if collisions_2380:
        for name, reasons in collisions_2380:
            a(f"  - `{name}`：{'；'.join(reasons)}")
    else:
        a("  - （无）")
    a("")
    a("## 5. 对应表核对")
    a("")
    a("比对报告文档自述：")
    a("")
    for k, v in stated.items():
        a(f"- {k}：{v}")
    a("")
    a(f"- 同义表实际数据行：**{len(synonyms)}**（标题写 49，统计总览写 50；以表格实际行数为准）")
    a(f"- 找不到表实际词条：**{len(missing)}**")
    if map_meta["syn_duplicate_olds"]:
        a(f"- 同义表旧名重复：{map_meta['syn_duplicate_olds']}")
    if map_meta["missing_duplicate"]:
        a(f"- 找不到表名字重复：{map_meta['missing_duplicate']}")
    a("")
    a(f"规则 B 命中的同义表旧名（新名是带括号词条，视为直接对应）：**{len(map_meta['rule_b'])}**（期望 8）→ {_yn(len(map_meta['rule_b']) == 8)}")
    lines.extend(_bullet_list([f"`{n}`" for n in map_meta["rule_b"]]))
    a(f"规则 C 命中的找不到表旧名（恰为拆出的名字）：**{len(map_meta['rule_c'])}**（期望 5）→ {_yn(len(map_meta['rule_c']) == 5)}")
    lines.extend(_bullet_list([f"`{n}`" for n in map_meta["rule_c"]]))
    a("")
    counts = map_meta["counts"]
    n_same = counts.get("一致", 0)
    n_rename = counts.get("改名", 0)
    n_delete = counts.get("删除", 0)
    n_other = sum(c for k, c in counts.items() if k not in ("一致", "改名", "删除"))
    a(f"3.5 Concept 节点数：**{concept_node_count}**；CSV 中不同旧名：**{len(map_rows)}**。")
    a("")
    a("| 处理方式 | 数量 | 期望 | 结论 |")
    a("|---|---:|---:|---|")
    a(f"| 一致 | {n_same} | 1010 | {_yn(n_same == 1010)} |")
    a(f"| 改名 | {n_rename} | 42 | {_yn(n_rename == 42)} |")
    a(f"| 删除 | {n_delete} | 43 | {_yn(n_delete == 43)} |")
    a(f"| 合计（三类） | {n_same + n_rename + n_delete} | 1095 | {_yn(n_same + n_rename + n_delete == 1095)} |")
    if n_other:
        a("")
        a(f"另有未归类 **{n_other}** 条，见 CSV 中处理方式为「未归类」的行。")
    a("")
    a(f"### 报告中出现、但 3.5 不存在的旧名：**{len(map_meta['report_only'])}**")
    a("")
    if map_meta["report_only"]:
        for src, old, new in map_meta["report_only"]:
            extra = f" → `{new}`" if new else ""
            a(f"- {src}：`{old}`{extra}")
    else:
        a("- （无）")
    a("")
    a(f"### 改名后的新名不在 2380 名单中：**{len(map_meta['renamed_not_in_2380'])}**")
    a("")
    if map_meta["renamed_not_in_2380"]:
        for r in map_meta["renamed_not_in_2380"]:
            a(f"- `{r['旧名']}` → `{r['新名']}`")
    else:
        a("- （无）")
    a("")
    a(f"### 删除项的旧名却出现在 2380 名单中：**{len(map_meta['deleted_but_in_2380'])}**")
    a("")
    if map_meta["deleted_but_in_2380"]:
        for r in map_meta["deleted_but_in_2380"]:
            a(f"- `{r['旧名']}`")
    else:
        a("- （无）")
    a("")
    a(f"### 两个或更多旧名对应到同一个新名：**{len(map_meta['collisions'])}** 个新名")
    a("")
    if map_meta["collisions"]:
        for new, olds in sorted(map_meta["collisions"].items()):
            a(f"- `{new}` ← {'、'.join(f'`{o}`' for o in olds)}")
    else:
        a("- （无）")
    a("")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_inventory(
    path: Path,
    info: dict[str, str],
    server_agent: str,
    data: dict,
    concept_anomalies: dict,
    scripture_anomalies: dict,
    concept_props: list[tuple[str, int, int]],
    migration: list[dict],
    v4: dict,
) -> None:
    lines: list[str] = []
    a = lines.append
    a("# 3.5 图谱盘点")
    a("")
    a("本报告由 `back_mic/backend/panai4/graph/inventory_35.py` 生成。对 3.5 只执行了 MATCH / SHOW，经 `execute_read` 提交。报告不含密码。")
    a("")
    a(f"- 连接地址：`{info['uri']}`")
    a(f"- 用户：`{info['user']}`")
    a(f"- 服务器：{server_agent}")
    a("")
    node_total = sum(int(r["cnt"]) for r in data["labels"])
    a("## 1. 节点标签")
    a("")
    a(f"节点总数：**{node_total}**")
    a("")
    a("| 标签组合 | 数量 |")
    a("|---|---:|")
    for row in data["labels"]:
        a(f"| {_label_key(row['labels'])} | {row['cnt']} |")
    a("")
    a("## 2. 关系类型")
    a("")
    rel_total = sum(int(r["cnt"]) for r in data["rels"])
    by_type: dict[str, int] = defaultdict(int)
    for row in data["rels"]:
        by_type[row["rel_type"]] += int(row["cnt"])
    a(f"关系总数：**{rel_total}**（有向，按起点→终点计数）")
    a("")
    a("| 关系类型 | 数量 |")
    a("|---|---:|")
    for name, cnt in sorted(by_type.items()):
        a(f"| {name} | {cnt} |")
    a("")
    a("| 关系类型 | 起点标签 | 终点标签 | 数量 |")
    a("|---|---|---|---:|")
    for row in data["rels"]:
        a(
            f"| {row['rel_type']} | {_label_key(row['start_labels'])} | "
            f"{_label_key(row['end_labels'])} | {row['cnt']} |"
        )
    a("")
    a("关系上的标注键：")
    a("")
    any_rel_prop = False
    for row in data["rel_keys"]:
        if row["prop"] is None:
            continue
        any_rel_prop = True
        a(f"- `{row['rel_type']}`.{row['prop']}：{row['cnt']} 条")
    if not any_rel_prop:
        a("- 所有关系都没有标注。")
    a("")
    a("## 3. Concept 与 Scripture 的标注键")
    a("")
    a("「有值」指不是 null、不是空字符串、不是空列表。")
    a("")
    a("### Concept")
    a("")
    a("| 标注 | 出现该键的节点数 | 其中有值的节点数 |")
    a("|---|---:|---:|")
    for key, present, nonempty in concept_props:
        a(f"| {key} | {present} | {nonempty} |")
    if not concept_props:
        a("| （无） | 0 | 0 |")
    a("")
    a("### Scripture")
    a("")
    a("Scripture 的正文不拉回本地。下表来自库内聚合。「有值」已排除 null 和空字符串。")
    a("")
    a("| 标注 | 出现该键的节点数 | 其中有值的节点数 |")
    a("|---|---:|---:|")
    for row in data["scripture_key_stats"]:
        present = int(row["present"])
        nonempty = int(row["nonempty_scalar"])
        a(f"| {row['prop']} | {present} | {nonempty} |")
    if not data["scripture_key_stats"]:
        a("| （无） | 0 | 0 |")
    a("")
    a("## 4. Concept.name")
    a("")
    a(f"- 节点数：**{concept_anomalies['total']}**")
    a(f"- 不同 name：**{concept_anomalies['distinct']}**")
    a(f"- name 为空（null）：**{concept_anomalies['nulls']}**")
    a(f"- name 为空白字符串：**{len(concept_anomalies['blanks'])}**")
    a(f"- 重复的 name：**{len(concept_anomalies['duplicates'])}** → {_yn(not concept_anomalies['duplicates'] and concept_anomalies['nulls'] == 0)}")
    if concept_anomalies["duplicates"]:
        lines.extend(_bullet_list([f"`{n}` ×{c}" for n, c in concept_anomalies["duplicates"]]))
    a(f"- 首尾有空格或全角空格：**{len(concept_anomalies['padded'])}**")
    lines.extend(_bullet_list([repr(x) for x in concept_anomalies["padded"]]))
    a(f"- 含半角括号 `()`：**{len(concept_anomalies['half'])}**")
    lines.extend(_bullet_list([f"`{x}`" for x in concept_anomalies["half"]]))
    a(f"- 含全角括号 `（）`：**{len(concept_anomalies['full'])}**")
    lines.extend(_bullet_list([f"`{x}`" for x in concept_anomalies["full"]]))
    a(f"- 全角与半角括号混用：**{len(concept_anomalies['mixed'])}**")
    lines.extend(_bullet_list([f"`{x}`" for x in concept_anomalies["mixed"]]))
    a(f"- 内部含连续空格或全角空格：**{len(concept_anomalies['inner_space'])}**")
    lines.extend(_bullet_list([repr(x) for x in concept_anomalies["inner_space"]]))
    a("")
    a("库内约束（SHOW CONSTRAINTS，只读）：")
    a("")
    if data["constraints"]:
        for row in data["constraints"]:
            name = row.get("name")
            ctype = row.get("type")
            entity = row.get("labelsOrTypes") or row.get("entityType")
            props = row.get("properties")
            a(f"- {name}：type={ctype}，entity={entity}，properties={props}")
    else:
        a("- 没有约束。name 的唯一性只来自数据，没有数据库约束。")
    a("")
    a("库内索引（SHOW INDEXES，只读）：")
    a("")
    if data["indexes"]:
        for row in data["indexes"]:
            a(
                f"- {row.get('name')}：type={row.get('type')}，"
                f"entity={row.get('entityType')}，labels={row.get('labelsOrTypes')}，"
                f"properties={row.get('properties')}"
            )
    else:
        a("- 没有索引。")
    a("")
    a("## 5. Scripture.id")
    a("")
    a(f"- 节点数：**{scripture_anomalies['total']}**")
    a(f"- 不同 id：**{scripture_anomalies['distinct']}**")
    a(f"- id 为空（null）：**{scripture_anomalies['nulls']}**")
    a(f"- id 为空白字符串：**{len(scripture_anomalies['blanks'])}**")
    a(f"- 重复的 id：**{len(scripture_anomalies['duplicates'])}** → {_yn(not scripture_anomalies['duplicates'] and scripture_anomalies['nulls'] == 0 and not scripture_anomalies['blanks'])}")
    if scripture_anomalies["duplicates"]:
        lines.extend(_bullet_list([f"`{n}` ×{c}" for n, c in scripture_anomalies["duplicates"]]))
    a(f"- 首尾有空格：**{len(scripture_anomalies['padded'])}**")
    lines.extend(_bullet_list([repr(x) for x in scripture_anomalies["padded"]]))
    a("")
    a("## 6. 迁移预估（只读，未写入）")
    a("")
    a("依据 `docs/panai4_graph/concept_mapping_35_to_40.csv`。")
    a("Concept 一端：处理方式为「一致」或「改名」视为能对应，新名用 CSV 的「新名」；「删除」和「未归类」不能对应。")
    a("Scripture 一端：全部保留，身份用 `id`（本库 Scripture 不在删除名单里）。")
    a("「会重复」指两端都对应上之后，同一关系类型、同一新起点、同一新终点出现多于一次。下表「会重复的关系数」是多出来的条数（每组保留 1 条后的余数），不是组数。")
    a("")
    a("| 关系类型 | 3.5 原有 | 两端都能对应 | 因某一端被删除而舍弃 | 无法判定 | 会重复的关系数 | 重复组数 |")
    a("|---|---:|---:|---:|---:|---:|---:|")
    for row in migration:
        a(
            f"| {row['rel_type']} | {row['total']} | {row['both']} | {row['discarded']} | "
            f"{row['unmapped']} | {row['extra']} | {row['dup_groups']} |"
        )
    a("")
    for row in migration:
        if row["examples"]:
            a(f"### {row['rel_type']} 的重复组（最多 15 组）")
            a("")
            for ex in row["examples"]:
                a(f"- {ex}")
            a("")
    a("## 7. neo4j-v4 连接测试")
    a("")
    if not v4.get("configured"):
        a("未配置。`NEO4J_V4_URI` / `NEO4J_V4_USER` / `NEO4J_V4_PASSWORD` 至少有一项在环境中为空或缺失，因此没有尝试连接，也没有写入。")
        missing = v4.get("missing") or []
        if missing:
            a("")
            a("缺失或为空的变量：")
            for name in missing:
                a(f"- `{name}`")
    else:
        a(f"- 连接地址：`{v4.get('uri')}`")
        a(f"- 用户：`{v4.get('user')}`")
        if v4.get("ok"):
            a(f"- 能否连上：**能**")
            a(f"- Neo4j 版本：{v4.get('agent')}")
            a(f"- 现有节点数：**{v4.get('nodes')}**")
            a(f"- 现有关系数：**{v4.get('rels')}**")
        else:
            a("- 能否连上：**不能**")
            a(f"- 错误：{v4.get('error')}")
    a("")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["旧名", "新名", "处理方式", "备注"])
        writer.writeheader()
        writer.writerows(rows)


def _collisions_2380(terms: list[str], splits: list[tuple[str, str, str]], list_2380: list[str]) -> list[tuple[str, list[str]]]:
    original = set(terms)
    reasons: dict[str, list[str]] = defaultdict(list)
    produced: dict[str, list[str]] = defaultdict(list)
    for src, outer, inner in splits:
        for part, side in ((outer, "括号外"), (inner, "括号内")):
            produced[part].append(f"{side}来自 `{src}`")
            if part in original:
                reasons[part].append(f"与原词表已有词条重名（拆自 `{src}` 的{side}）")
    for part, origins in produced.items():
        if len(origins) > 1:
            reasons[part].append("多个括号词条拆出了同一个名字：" + "；".join(origins))
    # 2380 名单内部重复（含非拆分造成的）
    counts = Counter(list_2380)
    for name, cnt in counts.items():
        if cnt > 1 and name not in reasons:
            reasons[name].append(f"在 2380 名单中出现 {cnt} 次")
    return sorted(reasons.items())


def probe_v4() -> dict:
    keys = ("NEO4J_V4_URI", "NEO4J_V4_USER", "NEO4J_V4_PASSWORD")
    missing = [k for k in keys if not (os.environ.get(k) or "").strip()]
    if missing:
        return {"configured": False, "missing": missing}
    uri = os.environ["NEO4J_V4_URI"].strip()
    user = os.environ["NEO4J_V4_USER"].strip()
    password = os.environ["NEO4J_V4_PASSWORD"]
    driver = None
    try:
        driver = GraphDatabase.driver(uri, auth=(user, password))
        driver.verify_connectivity()
        agent = str(driver.get_server_info().agent)
        with driver.session(default_access_mode=READ_ACCESS) as session:
            nodes = _run(session, "MATCH (n) RETURN count(n) AS c")[0]["c"]
            rels = _run(session, "MATCH ()-[r]->() RETURN count(r) AS c")[0]["c"]
        return {
            "configured": True,
            "ok": True,
            "uri": uri,
            "user": user,
            "agent": agent,
            "nodes": nodes,
            "rels": rels,
        }
    except Exception as exc:
        return {
            "configured": True,
            "ok": False,
            "uri": uri,
            "user": user,
            "error": _redact(str(exc), password),
        }
    finally:
        if driver is not None:
            driver.close()


def main() -> int:
    _REPORTS.mkdir(parents=True, exist_ok=True)
    preview = _REPORTS / "_docx_preview.txt"
    if preview.exists():
        preview.unlink()

    env = _load_env()
    terms, dirty_lines = load_terms_txt()
    term_dupes = sorted(n for n, c in Counter(terms).items() if c > 1)
    list_2380, splits, bad_parens = build_2380(terms)
    collisions = _collisions_2380(terms, splits, list_2380)
    trunk_text = _TRUNKS_TXT.read_text(encoding="utf-8-sig")
    trunk_pairs, trunk_notes = parse_trunks(trunk_text)
    term_rows = json.loads(_TERMS_JSON.read_text(encoding="utf-8-sig"))
    tag_obj = json.loads(_TAGS_JSON.read_text(encoding="utf-8-sig"))
    synonyms, missing, stated = load_docx()
    cross = cross_check(terms, term_rows, tag_obj, trunk_pairs)

    print(f"连接 3.5：{env['uri']} 用户 {env['user']}（密码不打印）")
    driver = GraphDatabase.driver(env["uri"], auth=(env["user"], env["password"]))
    try:
        driver.verify_connectivity()
        agent = str(driver.get_server_info().agent)
        with driver.session(default_access_mode=READ_ACCESS) as session:
            data = inventory_graph(session)
    finally:
        driver.close()
    print(f"3.5 已读取，服务器 {agent}，节点标签组 {len(data['labels'])}，关系行 {len(data['edges'])}")

    concept_props_raw = []
    concept_names = []
    for row in data["concepts"]:
        props = row["props"] or {}
        concept_props_raw.append(props)
        concept_names.append(props.get("name"))
    concept_prop_stats = _prop_stats(concept_props_raw)
    anomalies_input = []
    for n in concept_names:
        anomalies_input.append(n if n is None else str(n))
    concept_anomalies = _name_anomalies(anomalies_input)
    scripture_ids = [row["id"] for row in data["scriptures"]]
    scripture_anomalies = _name_anomalies([i if i is None else str(i) for i in scripture_ids])

    # 对应表以图中的 Concept.name 为准；null name 用空字符串占位并标未归类
    names_for_map = []
    for n in concept_names:
        names_for_map.append(n if isinstance(n, str) else "")
    map_rows, map_meta = build_mapping(
        names_for_map,
        set(list_2380),
        splits,
        synonyms,
        missing,
    )
    csv_path = _DATA / "concept_mapping_35_to_40.csv"
    write_csv(csv_path, map_rows)
    migration = migration_estimate(data["edges"], map_rows)
    v4 = probe_v4()

    write_source_check(
        _REPORTS / "source_check.md",
        terms,
        dirty_lines,
        term_dupes,
        splits,
        bad_parens,
        list_2380,
        collisions,
        trunk_pairs,
        trunk_notes,
        cross,
        synonyms,
        missing,
        stated,
        map_rows,
        map_meta,
        concept_node_count=len(concept_names),
    )
    write_inventory(
        _REPORTS / "35_inventory.md",
        {"uri": env["uri"], "user": env["user"]},
        agent,
        data,
        concept_anomalies,
        scripture_anomalies,
        concept_prop_stats,
        migration,
        v4,
    )
    counts = map_meta["counts"]
    print(
        "对应表",
        dict(counts),
        "规则B",
        len(map_meta["rule_b"]),
        "规则C",
        len(map_meta["rule_c"]),
    )
    print("已写", csv_path)
    print("已写", _REPORTS / "35_inventory.md")
    print("已写", _REPORTS / "source_check.md")
    return 0


if __name__ == "__main__":
    sys.exit(main())
