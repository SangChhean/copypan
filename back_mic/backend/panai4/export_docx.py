# -*- coding: utf-8 -*-
"""按步骤导出 docx。不写入运行编号、费用或 Prompt 版本。"""
from __future__ import annotations

import io
import json
import logging
import re
import zipfile
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .pipeline import STEP_LABELS, extract_burden_body

_log = logging.getLogger("panai4.export")
_KNOWN_TIERS = ("最内在", "内在", "外在")

EXPORT_STEPS = ("select", "burden", "skeleton", "orig_skeleton", "diagnosis")
# 旧流程 zip 仍只含这四步（无选词）
LEGACY_EXPORT_STEPS = ("burden", "skeleton", "orig_skeleton", "diagnosis")
FULL_ARTICLE_SECTIONS = (
    ("topic", "主题"),
    ("select", "选出的词条"),
    ("burden", "负担说明"),
    ("skeleton", "新生成的龙骨"),
    ("orig_skeleton", "原纲目的龙骨"),
    ("diagnosis", "对比诊断"),
)
_FONT = "Microsoft YaHei"
_ILLEGAL = set('\\/:*?"<>|')


def safe_filename_part(text: str, limit: int | None = None) -> str:
    raw = (text or "").strip()
    if limit is not None:
        raw = raw[:limit]
    cleaned = []
    for char in raw:
        if char in _ILLEGAL or ord(char) < 32:
            cleaned.append("_")
        else:
            cleaned.append(char)
    name = "".join(cleaned).strip(" .")
    return name or "未命名"


def docx_filename(position: int, title: str, step: str) -> str:
    label = STEP_LABELS.get(step, step)
    return f"【{label}】第{position}篇_{safe_filename_part(title, 30)}.docx"


def article_folder(position: int, title: str) -> str:
    return f"第{position}篇_{safe_filename_part(title, 30)}"


def content_disposition(filename: str) -> str:
    encoded = quote(filename)
    return f"attachment; filename*=UTF-8''{encoded}"


def _set_run_font(run, size_pt: float, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = _FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), _FONT)


def _set_paragraph_format(paragraph, *, align=None, heading: bool = False) -> None:
    fmt = paragraph.paragraph_format
    if heading:
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
        fmt.space_before = Pt(6)
        fmt.space_after = Pt(6)
    else:
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(20)
        fmt.space_before = Pt(2)
        fmt.space_after = Pt(2)
    if align is not None:
        paragraph.alignment = align


def _em_width(text: str) -> float:
    width = 0.0
    for char in text:
        width += 6.0 if ord(char) < 128 else 12.0
    return width


_HASH = re.compile(r"^(#{1,6})\s?")
_BULLET = re.compile(r"^[-*]\s+")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_NUMBERED = re.compile(r"^(\d+\.\s+)")
_PAREN_CN = re.compile(r"^([（(][一二三四五六七八九十]+[）)])")
_PAREN_NUM = re.compile(r"^([（(]\d+[）)])")
_LETTER = re.compile(r"^([A-Za-z])(?![A-Za-z])")
_MAJOR = re.compile(r"^([壹贰叁肆伍陆柒捌玖拾]+)")
_MINOR = re.compile(r"^([一二三四五六七八九十]+)")
_DIGIT = re.compile(r"^(\d+)")


def _bold_runs(text: str) -> list[tuple[str, bool]]:
    parts: list[tuple[str, bool]] = []
    pos = 0
    for match in _BOLD.finditer(text):
        if match.start() > pos:
            parts.append((text[pos : match.start()], False))
        parts.append((match.group(1), True))
        pos = match.end()
    if pos < len(text) or not parts:
        parts.append((text[pos:], False))
    return parts


def model_paragraphs(text: str, step: str) -> list[dict]:
    numbered = step in {"skeleton", "orig_skeleton"}
    number = 0
    paragraphs = []
    for line in _split_lines(text):
        heading = False
        hash_match = _HASH.match(line)
        if hash_match:
            line = line[hash_match.end() :]
            heading = True
        prefix = ""
        bullet = _BULLET.match(line)
        if bullet and not heading:
            line = line[bullet.end() :]
            if numbered:
                number += 1
                prefix = f"{number}. "
        elif line.strip():
            number = 0
            existing = _NUMBERED.match(line) if numbered else None
            if existing:
                prefix = existing.group(1)
                line = line[existing.end() :]
        runs = _bold_runs(line)
        if prefix:
            runs = [(prefix, False), *runs]
        if heading or (step == "diagnosis" and "".join(piece for piece, _bold in runs).startswith("【")):
            runs = [(piece, True) for piece, _bold in runs]
        paragraphs.append({"runs": runs, "hang": _em_width(prefix), "left": 0.0})
    return paragraphs


def format_outline_line(line: str) -> tuple[float, float, str]:
    """返回（层级左缩进磅值, 悬挂宽度磅值, 显示文字）。不删除、不改动非 Tab 的文字。"""
    if line == "":
        return 0.0, 0.0, ""
    for level, pattern in (
        (6, _PAREN_CN),
        (6, _PAREN_NUM),
        (6, _LETTER),
        (0, _MAJOR),
        (2, _MINOR),
        (4, _DIGIT),
    ):
        match = pattern.match(line)
        if not match:
            continue
        prefix = match.group(1)
        rest = line[match.end() :]
        if rest.startswith(("、", ".", "．")):
            prefix += rest[0]
            rest = rest[1:]
        if rest.startswith("\t"):
            rest = rest[1:]
            prefix += "\u3000"
        return float(level * 12), _em_width(prefix), prefix + rest
    return 0.0, 0.0, line


def _add_heading(document: Document, text: str, size_pt: float, align=None) -> None:
    paragraph = document.add_paragraph()
    _set_paragraph_format(paragraph, align=align, heading=True)
    run = paragraph.add_run(text)
    _set_run_font(run, size_pt, True)


def _add_body(document: Document, runs: list[tuple[str, bool]], *, left: float = 0, hang: float = 0) -> None:
    paragraph = document.add_paragraph()
    _set_paragraph_format(paragraph, heading=False)
    if hang or left:
        paragraph.paragraph_format.left_indent = Pt(left + hang)
        paragraph.paragraph_format.first_line_indent = Pt(-hang)
    if not runs:
        runs = [("", False)]
    for text, bold in runs:
        run = paragraph.add_run(text)
        _set_run_font(run, 12, bold)


def _prepare_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.gutter = Cm(0)
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.text = ""
    _add_page_number(footer)
    normal = document.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), _FONT)
    return document


_LIST_MARK = re.compile(r"^(?:[-*•]\s+|•|\d+[.、．)]\s*|\d+\s+)")


def number_skeleton(text: str) -> str:
    """只用于显示和导出。不写回数据库。"""
    lines = (text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    listed = any(line.strip() and _LIST_MARK.match(line) for line in lines)
    if listed:
        number = 0
        out = []
        for line in lines:
            if not line.strip() or _HASH.match(line):
                out.append(line)
                continue
            number += 1
            out.append(f"{number}. {_LIST_MARK.sub('', line, count=1)}")
        return "\n".join(out)
    paragraphs = []
    current = []
    for line in lines:
        if not line.strip():
            if current:
                paragraphs.append(current)
                current = []
        else:
            current.append(line)
    if current:
        paragraphs.append(current)
    out = []
    for index, paragraph in enumerate(paragraphs, start=1):
        if out:
            out.append("")
        first, *rest = paragraph
        out.append(f"{index}. {first}")
        out.extend(rest)
    return "\n".join(out)


def build_docx(title: str, step: str, body: str, appendix: str | None = None, position: int = 1) -> bytes:
    document = _prepare_document()
    if step in {"skeleton", "orig_skeleton"}:
        body = number_skeleton(body)
    _add_heading(document, f"第{position}篇\u3000{title}", 16, WD_ALIGN_PARAGRAPH.CENTER)
    _add_heading(document, STEP_LABELS.get(step, step), 14)
    for paragraph in model_paragraphs(body, step):
        _add_body(document, paragraph["runs"], left=paragraph["left"], hang=paragraph["hang"])
    if appendix is not None:
        _add_heading(document, "附：原纲目全文", 14)
        for line in _split_lines(appendix):
            left, hang, text = format_outline_line(line)
            _add_body(document, [(text, False)], left=left, hang=hang)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _holder():
        run = paragraph.add_run()
        _set_run_font(run, 12, False)
        return run._r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    _holder().append(begin)
    _holder().append(instr)
    tail = _holder()
    tail.append(separate)
    tail.append(end)


def _split_lines(text: str) -> list[str]:
    normalized = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    if normalized == "":
        return [""]
    return normalized.split("\n")


def step_body(
    step: str,
    output_text: str,
    original_outline: str | None,
    *,
    flow_kind: str = "legacy",
) -> tuple[str, str | None]:
    if step == "burden":
        if flow_kind == "full":
            from .sc.export_sc import extract_burden_export

            body, _warning = extract_burden_export(output_text or "")
            return body, None
        return extract_burden_body(output_text or ""), None
    if step == "orig_skeleton":
        return output_text or "", original_outline or ""
    return output_text or "", None


def _parse_json_field(raw) -> list | dict | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, (list, dict)):
        return raw
    try:
        return json.loads(raw)
    except (TypeError, json.JSONDecodeError):
        return None


def _tier_lookup() -> dict[str, str]:
    """词条名 → 档次（来自词条三档分类.json）。加载失败时返回空表。"""
    try:
        from panai4.sc.data import load_sc_data

        return dict(load_sc_data().tier_of or {})
    except Exception as exc:  # noqa: BLE001 — 导出不应因词表加载失败而崩
        _log.warning("选词导出：加载三档分类失败，仅用已有档次字段：%s", exc)
        return {}


def _resolve_tier(name: str, raw_tier: str | None, lookup: dict[str, str]) -> str:
    """档次优先用已有合法值；否则按词条名回查三档分类；再没有则「未分档」。"""
    tier = (raw_tier or "").strip()
    if tier in _KNOWN_TIERS:
        return tier
    looked = (lookup.get(name) or "").strip()
    if looked in _KNOWN_TIERS:
        return looked
    return "未分档"


def final_terms_for_export(item: dict, steps: list[dict] | None = None) -> list[dict]:
    """最终词表：优先 final_terms_json；否则 select.meta_json.items。

    档次：合法「最内在/内在/外在」用原值；否则按词条名回查三档分类；
    回查不到标「未分档」（不默认外在）。不改写数据库。
    """
    final = _parse_json_field(item.get("final_terms_json"))
    if isinstance(final, list) and final:
        rows = final
    else:
        rows = []
        for step in steps or []:
            if step.get("step") != "select":
                continue
            meta = _parse_json_field(step.get("meta_json")) or {}
            rows = list(meta.get("items") or meta.get("candidates") or [])
            break
    select_meta = {}
    for step in steps or []:
        if step.get("step") == "select":
            select_meta = _parse_json_field(step.get("meta_json")) or {}
            break
    original = {
        (row.get("name") or "").strip(): row
        for row in (select_meta.get("items") or select_meta.get("candidates") or [])
        if (row.get("name") or "").strip()
    }
    lookup = _tier_lookup()
    out = []
    for row in rows:
        name = (row.get("name") or "").strip()
        if not name:
            continue
        raw_tier = row.get("tier") or row.get("tier_actual") or ""
        tier = _resolve_tier(name, raw_tier, lookup)
        reason = row.get("reason") or ""
        src = original.get(name)
        adjusted = False
        if name not in original:
            adjusted = True
        elif src and (src.get("tier") or src.get("tier_actual") or "") != (raw_tier or tier):
            adjusted = True
        out.append(
            {
                "name": name,
                "tier": tier,
                "reason": reason,
                "manual_adjusted": adjusted,
            }
        )
    if original and set(original) != {r["name"] for r in out}:
        for row in out:
            if row["name"] not in original:
                row["manual_adjusted"] = True
    return out


def format_select_export_text(terms: list[dict]) -> str:
    """按档次分组输出。注意：不可用「空列表 or 外在」——空列表在 Python 为假，会全进外在。"""
    by_tier: dict[str, list[dict]] = {t: [] for t in _KNOWN_TIERS}
    by_tier["未分档"] = []
    for row in terms:
        tier = row.get("tier") or "未分档"
        if tier not in by_tier:
            tier = "未分档"
        by_tier[tier].append(row)
    unknown = by_tier["未分档"]
    if unknown:
        names = "、".join(r["name"] for r in unknown[:20])
        more = f" 等{len(unknown)}条" if len(unknown) > 20 else ""
        _log.warning("选词导出：%d 条词条未分档，列入【未分档】：%s%s", len(unknown), names, more)
    lines: list[str] = []
    # 三档标题始终输出（含 0 条），便于核对；未分档仅在有条目时出现
    for tier in _KNOWN_TIERS:
        rows = by_tier[tier]
        lines.append(f"【{tier}】（{len(rows)}）")
        for row in rows:
            mark = "（人工调整）" if row.get("manual_adjusted") else ""
            reason = row.get("reason") or "（无依据）"
            lines.append(f"{row['name']}{mark}｜{reason}")
        lines.append("")
    rows_u = by_tier["未分档"]
    if rows_u:
        lines.append(f"【未分档】（{len(rows_u)}）")
        for row in rows_u:
            mark = "（人工调整）" if row.get("manual_adjusted") else ""
            reason = row.get("reason") or "（无依据）"
            lines.append(f"{row['name']}{mark}｜{reason}")
        lines.append("")
    return "\n".join(lines).rstrip() + ("\n" if lines else "")


def build_select_docx(item: dict, steps: list[dict], *, flow_kind: str = "full") -> tuple[str, bytes]:
    terms = final_terms_for_export(item, steps)
    body = format_select_export_text(terms)
    filename = docx_filename(item["position"], item["title"], "select")
    document = _prepare_document()
    _add_heading(document, f"第{item['position']}篇\u3000{item['title']}", 16, WD_ALIGN_PARAGRAPH.CENTER)
    _add_heading(document, "选出的词条", 14)
    for line in _split_lines(body):
        _add_body(document, [(line, False)])
    buffer = io.BytesIO()
    document.save(buffer)
    return filename, buffer.getvalue()


def build_full_article_docx(item: dict, steps: list[dict]) -> tuple[str, bytes]:
    """全流程一篇合并 docx：主题→词条→负担→龙骨→原纲目龙骨→对比诊断。"""
    by_step = {s["step"]: s for s in steps if s.get("status") == "done"}
    document = _prepare_document()
    title = item.get("title") or ""
    position = item.get("position") or 1
    _add_heading(document, f"第{position}篇\u3000{title}", 16, WD_ALIGN_PARAGRAPH.CENTER)

    # 主题
    _add_heading(document, "主题", 14)
    _add_body(document, [(title, False)])

    # 选出的词条
    _add_heading(document, "选出的词条", 14)
    terms = final_terms_for_export(item, steps)
    select_text = format_select_export_text(terms) if terms else "（无）"
    for line in _split_lines(select_text):
        _add_body(document, [(line, False)])

    for step_id, label in (
        ("burden", "负担说明"),
        ("skeleton", "新生成的龙骨"),
        ("orig_skeleton", "原纲目的龙骨"),
        ("diagnosis", "对比诊断"),
    ):
        row = by_step.get(step_id)
        _add_heading(document, label, 14)
        if not row:
            _add_body(document, [("（本步未完成或未选择）", False)])
            continue
        body, appendix = step_body(
            step_id,
            row.get("output_text") or "",
            item.get("original_outline"),
            flow_kind="full",
        )
        if step_id in {"skeleton", "orig_skeleton"}:
            body = number_skeleton(body)
        for paragraph in model_paragraphs(body, step_id):
            _add_body(document, paragraph["runs"], left=paragraph["left"], hang=paragraph["hang"])
        if appendix:
            _add_heading(document, "附：原纲目全文", 14)
            for line in _split_lines(appendix):
                left, hang, text = format_outline_line(line)
                _add_body(document, [(text, False)], left=left, hang=hang)

    filename = f"【全流程】第{position}篇_{safe_filename_part(title, 30)}.docx"
    buffer = io.BytesIO()
    document.save(buffer)
    return filename, buffer.getvalue()


def build_step_docx(
    item: dict, step_row: dict, *, flow_kind: str = "legacy", steps: list[dict] | None = None
) -> tuple[str, bytes]:
    step = step_row["step"]
    if step == "select":
        return build_select_docx(item, steps or [step_row], flow_kind=flow_kind)
    body, appendix = step_body(
        step,
        step_row.get("output_text") or "",
        item.get("original_outline"),
        flow_kind=flow_kind,
    )
    filename = docx_filename(item["position"], item["title"], step)
    return filename, build_docx(item["title"], step, body, appendix, item["position"])


def completed_export_steps(steps: list[dict], *, flow_kind: str = "legacy") -> list[dict]:
    allowed = EXPORT_STEPS if flow_kind == "full" else LEGACY_EXPORT_STEPS
    order = {name: index for index, name in enumerate(allowed)}
    rows = [step for step in steps if step.get("step") in allowed and step.get("status") == "done"]
    # 全流程：有 final 词表即可导出选词（即使 await_confirm 的展示不同）
    if flow_kind == "full":
        has_select = any(s.get("step") == "select" for s in rows)
        if not has_select:
            # 若 select done 但未进 rows——已覆盖；若只有 final_terms，由 zip 层补
            pass
    rows.sort(key=lambda row: order.get(row.get("step"), 99))
    return rows


def build_article_zip(
    item: dict, steps: list[dict], *, flow_kind: str = "legacy"
) -> tuple[str, bytes]:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        if flow_kind == "full":
            combined_name, combined = build_full_article_docx(item, steps)
            archive.writestr(combined_name, combined)
            # 选词单独一份（有最终词表或 select done）
            if item.get("final_terms_json") or any(
                s.get("step") == "select" and s.get("status") == "done" for s in steps
            ):
                sel_name, sel_data = build_select_docx(item, steps, flow_kind=flow_kind)
                archive.writestr(sel_name, sel_data)
        for step in completed_export_steps(steps, flow_kind=flow_kind):
            if step.get("step") == "select" and flow_kind == "full":
                continue  # 已单独写入
            name, data = build_step_docx(item, step, flow_kind=flow_kind, steps=steps)
            archive.writestr(name, data)
    filename = f"{article_folder(item['position'], item['title'])}.zip"
    return filename, buffer.getvalue()


def build_run_zip(
    items: list[tuple[dict, list[dict]]], *, flow_kind: str = "legacy"
) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for item, steps in items:
            folder = article_folder(item["position"], item["title"])
            if flow_kind == "full":
                combined_name, combined = build_full_article_docx(item, steps)
                archive.writestr(f"{folder}/{combined_name}", combined)
                if item.get("final_terms_json") or any(
                    s.get("step") == "select" and s.get("status") == "done" for s in steps
                ):
                    sel_name, sel_data = build_select_docx(item, steps, flow_kind=flow_kind)
                    archive.writestr(f"{folder}/{sel_name}", sel_data)
            for step in completed_export_steps(steps, flow_kind=flow_kind):
                if step.get("step") == "select" and flow_kind == "full":
                    continue
                name, data = build_step_docx(item, step, flow_kind=flow_kind, steps=steps)
                archive.writestr(f"{folder}/{name}", data)
    return buffer.getvalue()
